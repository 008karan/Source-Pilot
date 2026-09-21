"""Native-first extraction and deterministic, evidence-backed business projections.

No function in this module reads demo_state.json or normalized_truth.csv.
"""
from __future__ import annotations
import copy
import email
import hashlib
import json
import math
import re
from pathlib import Path
from uuid import uuid4
import pdfplumber
from docx import Document
from openpyxl import load_workbook
from . import repository as repo
from .ai import ai_extract_file, ai_extract_text, has_ai, judge_compliance, MODEL
from .intake_docs import email_lines
from .native_extractors import extract_native
from .store import DATA_DIR

def source_blocks(path):
    suffix = path.suffix.lower()
    if suffix == '.xlsx':
        wb = load_workbook(path, data_only=True, read_only=True)
        blocks = [{'sheet':s.title, 'row':i, 'text':' | '.join(str(x) for x in row if x is not None)}
                  for s in wb for i,row in enumerate(s.values,1) if any(x is not None for x in row)]
        wb.close()
        return blocks
    if suffix == '.docx':
        doc = Document(path)
        return ([{'paragraph':i,'text':p.text} for i,p in enumerate(doc.paragraphs) if p.text] +
                [{'table':t,'row':i,'text':' | '.join(c.text for c in row.cells)}
                 for t,table in enumerate(doc.tables) for i,row in enumerate(table.rows,1)])
    if suffix == '.pdf':
        with pdfplumber.open(path) as pdf:
            return [{'page':i,'text':p.extract_text() or ''} for i,p in enumerate(pdf.pages,1)]
    if suffix == '.txt':
        body = path.read_text()
        return [{'body_span':[m.start(),m.end()],'text':m.group()} for m in re.finditer(r'[^\n]+',body)]
    if suffix == '.eml':
        # Suppliers routinely send the price table as the HTML alternative only.
        return [{'line':i,'text':line} for i,line in enumerate(email_lines(path),1)]
    return []

# Wording suppliers actually use for "our price already gets it to your door".
CARRIAGE=r'(?:freight|delivery|shipping|transport(?:ation)?|carriage)'
FREIGHT_INCLUDED=re.compile(
    CARRIAGE+r'[\s:|,-]*(?:charges?[\s:|,-]*)?(?:are|is|be)?\s*(?:included|inclusive|free|prepaid|pre-paid|paid\s+by\s+(?:us|supplier|seller))'
    r'|(?:price|prices|rate|rates|quote|quotes|quoted|amount|amounts)[^.\n]{0,60}?(?:inclusive\s+of|includes?|including)\s+[^.\n]{0,40}?'+CARRIAGE+
    r'|inclusive\s+of\s+[^.\n]{0,30}?'+CARRIAGE+
    r'|\bDDP\b|\bCIP\b|\bCIF\b|free\s+delivery|door\s+delivery|landed\s+(?:price|cost|rate)s?'
    r'|\bF\.?\s?O\.?\s?R\.?\s+destination\b|delivered\s+(?:to\s+)?(?:site|destination|your\s+\w+)',re.I)
# ...and for "freight is your problem", which is a commercial fact, not a price.
FREIGHT_EXTRA=re.compile(
    CARRIAGE+r'[\s:|,-]*(?:charges?[\s:|,-]*)?(?:are|is|be|will\s+be)?\s*(?:extra|additional|excluded|not\s+included|at\s+actuals?|on\s+actuals?|to\s+(?:be\s+borne\s+by\s+)?(?:the\s+)?buyer)'
    r'|(?:to\s+)?(?:be\s+)?borne\s+by\s+(?:the\s+)?buyer'
    r'|\bex[\s-]?works\b|\bEXW\b|\bF\.?\s?O\.?\s?B\.?\b|\bF\.?\s?O\.?\s?R\.?\s+(?:works|factory)\b',re.I)

CURRENCY_MARKS={'INR':re.compile(r'₹|\bINR\b|\bRs\.?\b|\brupees?\b',re.I),
                'USD':re.compile(r'\bUSD\b|\bUS\$|\$\s?\d|\bdollars?\b',re.I)}

def consolidate_currency(terms,facts,rfx):
    """One quoting currency per response: from the lines, the document, or the RFx."""
    stated={f.get('raw_currency') for f in facts if f.get('raw_currency')}
    if len(stated)==1:
        return dict(terms,currency=stated.pop(),currency_source='stated on the quoted lines')
    if len(stated)>1:
        return dict(terms,currency='mixed',currency_source='different currencies appear on different lines')
    marked={code for code,pattern in CURRENCY_MARKS.items() if pattern.search(terms.get('source_text',''))}
    if len(marked)==1:
        return dict(terms,currency=marked.pop(),currency_source='stated in the response')
    if len(marked)>1:
        return dict(terms,currency='mixed',currency_source='more than one currency appears in the response')
    return dict(terms,currency=rfx.get('currency') or 'INR',currency_source='not stated; using the RFx quoting currency')

def terms_from_text(text):
    terms = {'source_text':text, 'freight':'unknown'}
    included=FREIGHT_INCLUDED.search(text);extra=FREIGHT_EXTRA.search(text)
    if included and not extra:
        terms.update(freight='included',freight_source=included.group().strip()[:200])
    elif extra:
        # A stated exclusion without a number stays unpriced until the buyer decides.
        terms.update(freight='excluded',freight_source=extra.group().strip()[:200])
    flat = re.search(CARRIAGE+r'[^.\n]{0,40}?(?:₹|INR|Rs\.?)\s*([\d.,]+)\s*(?:per|/)\s*(?:piece|pc|unit|each|item)',text,re.I)
    pct = re.search(CARRIAGE+r'[^.\n]{0,40}?([\d.]+)\s*%',text,re.I)
    # A freight amount with no per-unit qualifier is a charge for the whole order,
    # whether or not the supplier bothered to write the word "total".
    order = re.search(CARRIAGE+r'[^.\n;]{0,40}?(?P<cur>₹|INR|Rs\.?|USD|US\$|\$)\s*(?P<amount>\d[\d,]*(?:\.\d+)?)',
                      text,re.I)
    if flat:
        terms.update(freight='flat', freight_per_piece_inr=float(flat.group(1).replace(',','')),freight_source=flat.group().strip()[:200])
    elif pct:
        terms.update(freight='percent',freight_percent=float(pct.group(1))/100,freight_source=pct.group().strip()[:200])
    elif order and terms.get('freight') in (None,'unknown','excluded'):
        # Only when the response has not already said freight is included.
        terms.update(freight='order_total',freight_total=float(order.group('amount').replace(',','')),
                     freight_currency='USD' if order.group('cur').upper() in {'USD','US$','$'} else 'INR',
                     freight_source=order.group().strip()[:200])
    rebate = re.search(r'(\d+(?:\.\d+)?)% year-end volume rebate.*?USD\s*([\d,]+)',text,re.I|re.S)
    if rebate:
        terms['rebate']={'percent':float(rebate.group(1)),'threshold_usd':float(rebate.group(2).replace(',','')),'source_text':rebate.group()}
    return terms

def apply_terms(terms,answers,captured,rfx):
    """Supplier answers and captured terms are evidence too, not just prose."""
    if terms.get('freight') in {'included','flat','percent'}:return terms
    from .knowledge import question_kind
    qid=next((q['id'] for q in rfx['questionnaire'] if q['type']=='yes_no' and question_kind(q)=='freight'),None)
    answer=str((answers.get(qid) or {}).get('answer','')) if qid else ''
    if re.match(r'\s*(?:yes|true|included|inclusive)',answer,re.I):
        return dict(terms,freight='included',freight_source=f'{qid}: {answer}'[:200])
    for entry in captured:
        text=f"{entry.get('term_type','')}: {entry.get('value','')}"
        if not re.search(CARRIAGE,text,re.I):continue
        derived=terms_from_text(text)
        if derived['freight']!='unknown':
            return dict(terms,**{k:v for k,v in derived.items() if k!='source_text'})
    if re.match(r'\s*(?:no|false)',answer,re.I):
        return dict(terms,freight='excluded',freight_source=f'{qid}: {answer}'[:200])
    return terms

def native_answers(blocks, rfx):
    answers = {}
    for b in blocks:
        if b.get('sheet') == 'Quality':
            for q in rfx['questionnaire']:
                if b['text'].startswith(q['question']+' | '):
                    answers[q['id']]={'answer':b['text'].split(' | ')[-1], 'confidence':.99,'evidence':b}
    text = '\n'.join(b['text'] for b in blocks if b.get('sheet')!='Commercial Offer')
    patterns = {
        'Q1':r'ISO 9001\s+(?:valid|yes|certified)',
        'Q2':r'(?:FSC CoC\s+(?:valid|yes)|hold FSC Chain-of-Custody certification)',
        'Q3':r'(?:BF specs yes|meet stated burst factor|confirm compliance with the minimum burst-factor requirements)',
        'Q6':r'lead time.{0,25}?(\d+)\s*(?:calendar )?days',
    }
    for qid,pattern in patterns.items():
        if qid in answers:
            continue
        for b in blocks:
            m = re.search(pattern,b['text'],re.I)
            if m:
                answers[qid]={'answer':m.group(1) if qid=='Q6' else 'yes','confidence':.99,'evidence':b}
                break
    return answers

# A supplier saying plainly that something is missing is a failure, not a mystery.
NOT_MET = re.compile(r'non[-\s]?compliant|not compliant|not qualified|does not (?:comply|meet|hold)|'
                     r'no valid|not enclosed|not attached|not provided|not available|unable to (?:comply|provide)|'
                     r'expired|lapsed|cannot comply|^no\b|^false\b', re.I)
# ...and a list of valid certificates is a confirmation, whether or not it starts with "yes".
MET = re.compile(r'^(?:yes|true|confirmed|agreed|complied?)\b|\bvalid\b|\bcompliant\b|\bconfirmed\b|'
                 r'\bmeets?\b|\benclosed\b|\battached\b|\bregistered\b|\bcertified\b', re.I)
IN_PROGRESS = re.compile(r'in progress|under (?:renewal|review|process)|awaiting|applied for|renewal (?:application|is)|'
                         r'to be (?:provided|submitted)|pending', re.I)

def read_verdict(answer):
    """The deterministic reading, used offline and whenever the model cannot be reached."""
    text = str(answer or '').strip()
    if not text:
        return 'pending', 'No answer given'
    if NOT_MET.search(text):
        return 'fail', 'Supplier states a mandatory item is not met'
    if IN_PROGRESS.search(text):
        return 'pending', 'Supplier states an item is still in progress'
    if MET.search(text):
        return 'pass', 'Mandatory items confirmed'
    return 'pending', 'Answer does not clearly confirm the requirement'

RANK = {'pass': 0, 'pending': 1, 'fail': 2}

def assess(answers, text, rfx, use_ai=True):
    """Qualification and the reason for it, read from what the supplier actually wrote."""
    worst, why = 'pass', 'Mandatory declarations confirmed'
    for q in rfx['questionnaire']:
        if not q['mandatory']:
            continue
        given = answers.get(q['id'], {})
        answer = str(given.get('answer', '')).strip()
        if q['type'] == 'number':
            match = re.search(r'\d+(?:\.\d+)?', answer)
            if not match:
                verdict, reason = 'pending', f"{q['id']} has no stated figure"
            elif ('maximum' in q and float(match.group()) > q['maximum']) or \
                 ('minimum' in q and float(match.group()) < q['minimum']):
                verdict, reason = 'fail', f"{q['id']} is outside the stated limit"
            else:
                verdict, reason = 'pass', ''
        else:
            judged = judge_compliance(q['question'], answer) if (use_ai and answer) else None
            verdict, reason = (judged['verdict'], judged['reason']) if judged else read_verdict(answer)
        if given.get('confidence', 1) < .85 and verdict == 'pass':
            verdict, reason = 'pending', f"{q['id']} answer was read with low confidence"
        if RANK[verdict] > RANK[worst]:
            worst, why = verdict, reason or why
    if worst == 'pass' and re.search(r'renewal is in process|expires .*before', text, re.I):
        return 'pending', 'A certificate renewal is still in progress'
    return worst, why

def qualification(answers, text, rfx, use_ai=True):
    return assess(answers, text, rfx, use_ai)[0]

def context(rfx):
    return json.dumps({'items':[{k:v for k,v in i.items() if k!='baseline_unit_price_inr'} for i in rfx['items']], 'questionnaire':rfx['questionnaire']})

def extract(vendor_id, path, rfx, use_ai=True, native_fixture=False, progress=None):
    if progress:progress('parsing','Reading native structure and source locations')
    blocks = source_blocks(path)
    text = '\n'.join(b['text'] for b in blocks)
    result = None
    if native_fixture and path.suffix.lower() not in {'.jpg','.jpeg','.png','.webp'}:
        result = extract_native(vendor_id, path)
    if result is not None:
        facts = result['facts']
        answers = native_answers(blocks,rfx)
        method=result['method']
        unresolved=[]
        model=None
        captured=[]
    elif use_ai and has_ai():
        if progress:progress('extracting','Interpreting source facts and ambiguity')
        ai = ai_extract_text(text,context(rfx)) if blocks and path.suffix.lower()!='.pdf' else ai_extract_file(path.name,context(rfx),path)
        facts=ai['quote_lines']
        answers={q['question_id']:q for q in ai['questionnaire_answers']}
        method='gemini_vision' if not blocks else 'gemini_native_blocks'
        unresolved=ai['unresolved']
        model=MODEL
        captured=ai['commercial_terms']
        if not text:
            text='\n'.join(t['source_text'] for t in ai['commercial_terms'])
    else:
        return {'facts':[], 'terms':{}, 'questionnaire':{}, 'quality':'pending', 'method':'vision_required' if not blocks else 'ai_required', 'unresolved':['Runtime API key is required to interpret this document.'], 'blocks':blocks, 'model':None, 'captured_terms':[]}
    terms=terms_from_text(text)
    if progress:progress('validating','Checking line mapping, currencies, units and qualification')
    seen=set()
    for fact in facts:
        line=fact.get('line_no')
        if not isinstance(line,int) or not 1<=line<=len(rfx['items']) or line in seen:
            raise ValueError('Extraction returned an invalid or duplicate RFx line. Review source mapping.')
        seen.add(line)
        if fact.get('sku') != rfx['items'][line-1]['sku']:
            fact['quote_status']='unclear'
            fact['mapping_issue']='Source SKU does not match this RFx line'
        fact.setdefault('quote_status','quoted')
        fact.setdefault('confidence',.99)
        fact.setdefault('evidence',{'file':path.name,'source_text':fact.get('source_text','')})
        fact['extraction_method']=method
    quality,why=assess(answers,text,rfx,use_ai)
    return {'facts':facts,'terms':apply_terms(terms,answers,captured,rfx),'questionnaire':answers,'quality':quality,'quality_reason':why, 'method':method,'unresolved':unresolved,'blocks':blocks,'model':model,'captured_terms':captured}

def receive(vid,documents,message='',subject='Supplier response'):
    if vid not in repo.read()['vendors']:raise KeyError('Unknown supplier')
    if repo.read()['rfx_status']!='approved':raise ValueError('Approve the RFx before receiving responses.')
    def mutate(d):
        v=d['vendors'][vid]
        v.update(documents=documents,message=message,subject=subject,response_status='received',received_at=repo.now())
        # Old prices must not survive a newly received, not-yet-extracted offer.
        v.update(facts=[],terms={},questionnaire={},quality='pending',method='awaiting_processing',checksum=None)
        d['exceptions']=[e for e in d['exceptions'] if e['vendor_id']!=vid]
        d['stage']='collecting_responses'
        rebuild(d)
    return repo.change('response_received:'+vid,'supplier_inbox',mutate)

def receive_demo():
    missing={vid:record for vid,record in repo.demo_vendors().items() if vid not in repo.read()['vendors']}
    if missing:repo.change('demo_suppliers_added','buyer',lambda d:d['vendors'].update(copy.deepcopy(missing)))
    for vid,(_,filename) in repo.VENDORS.items():
        path=DATA_DIR/filename
        receive(vid,[{'path':str(path),'filename':filename,'fixture':True}],subject='FY27 quote — demo supplier message')
    return repo.read()

def ingest_received(vid,use_ai=True,job_id=None):
    from .jobs import progress as notify
    progress=lambda stage,message:notify(job_id,vid,stage,message)
    d=repo.read();v=d['vendors'][vid];docs=v.get('documents',[])
    if not docs:
        progress('waiting','No supplier response received yet');return d
    combined_digest=hashlib.sha256(''.join(hashlib.sha256(Path(x['path']).read_bytes()).hexdigest() for x in docs).encode()).hexdigest()
    if v.get('packet_checksum')==combined_digest and v['response_status']=='processed' and v.get('method') not in {'vision_required','ai_required'}:
        progress('completed','Previously processed source is unchanged');return d
    progress('received',f'Received {len(docs)} source document(s)')
    results=[]
    for doc in docs:
        path=Path(doc['path'])
        fixture_path=DATA_DIR/repo.VENDORS[vid][1] if vid in repo.VENDORS else None
        exact_fixture=bool(fixture_path and path.suffix==fixture_path.suffix and len(d['rfx']['items'])==30 and hashlib.sha256(path.read_bytes()).digest()==hashlib.sha256(fixture_path.read_bytes()).digest())
        # Message-only cover notes add evidence/terms; the attached document supplies quotes.
        if doc.get('cover_note') and len(docs)>1 and not (use_ai and has_ai()):
            blocks=source_blocks(path);text='\n'.join(b['text'] for b in blocks)
            result={'facts':[],'terms':terms_from_text(text),'questionnaire':native_answers(blocks,d['rfx']),'quality':'pending','method':'native_email_cover','unresolved':[],'blocks':blocks,'model':None,'captured_terms':[]}
        else:result=extract(vid,path,d['rfx'],use_ai,native_fixture=exact_fixture,progress=progress)
        for f in result['facts']:
            f.setdefault('evidence',{})['source_path']=str(path)
            f['evidence']['file']=doc['filename']
        result['source_document']=doc
        results.append(result)
    priced=lambda f:f.get('quote_status')=='quoted' and isinstance(f.get('raw_price'),(int,float)) and not isinstance(f.get('raw_price'),bool)
    main=next((r for r in results if any(priced(f) for f in r['facts'])),next((r for r in results if r['facts']),results[0]))
    merged=copy.deepcopy(main);facts={}
    for result in results:
        for f in result['facts']:
            current=facts.get(f['line_no'])
            if current is None or (priced(f) and not priced(current)):
                # The certificate attached beside a proposal has no price to give.
                facts[f['line_no']]=f
            elif priced(f) and priced(current) and any(f.get(k)!=current.get(k) for k in ['raw_price','raw_basis','raw_currency','pack_size','covers_units','covers_period_months']):
                current['quote_status']='unclear'
                current['mapping_issue']='Conflicting quotes across response attachments'
    merged['facts']=list(facts.values())
    alltext='\n'.join(r['terms'].get('source_text','') for r in results)
    merged['questionnaire']={k:v for r in results for k,v in r['questionnaire'].items()}
    captured=[];seen=set()
    for term in (t for r in results for t in r.get('captured_terms',[])):
        # The same cover note attached twice must not state its terms twice.
        key=(str(term.get('term_type','')).strip().lower(),str(term.get('value','')).strip().lower())
        if key in seen:continue
        seen.add(key);captured.append(term)
    merged['terms']=dict(apply_terms(terms_from_text(alltext),merged['questionnaire'],captured,d['rfx']),captured=captured)
    merged['quality'],merged['quality_reason']=assess(merged['questionnaire'],alltext,d['rfx'],use_ai)
    merged['unresolved']=list(dict.fromkeys(f"{r['source_document']['filename']}: {x}" for r in results for x in r['unresolved']))
    progress('normalizing','Converting prices, freight and currency with explicit rules')
    result=commit_extraction(vid,Path(main['source_document']['path']),merged)
    result=repo.change('response_packet_processed:'+vid,'workflow',lambda d:d['vendors'][vid].update(packet_checksum=combined_digest))
    progress('completed',f"{len(merged['facts'])} line attempts · qualification {merged['quality']}")
    return result

# "60-month rental term" on the requirement is the term, wherever the buyer wrote it.
TERM_WORDS=r'rental|lease|leasing|subscription|contract|tenure|term|licen[cs]e|service|support|amc|hire'
TERM_PATTERN=re.compile(
    r'(?:(?P<a>\d+(?:\.\d+)?)\s*-?\s*(?P<ua>month|months|year|years|yr|yrs)\b[^.\n]{0,28}?(?:'+TERM_WORDS+r')'
    r'|(?:'+TERM_WORDS+r')[^.\n]{0,28}?(?P<b>\d+(?:\.\d+)?)\s*-?\s*(?P<ub>month|months|year|years|yr|yrs)\b)',re.I)

def requirement_term(item,rfx):
    """Months the requirement runs for: the stated field, else the buyer's own wording."""
    stated=item.get('term_months')
    if isinstance(stated,(int,float)) and not isinstance(stated,bool) and math.isfinite(stated) and stated>0:
        return float(stated),'stated on the requirement line'
    for text,source in ((item.get('description'),'the requirement line'),(rfx.get('scope'),'the RFx scope')):
        match=TERM_PATTERN.search(text or '')
        if not match:continue
        value=float(match.group('a') or match.group('b'))
        if (match.group('ua') or match.group('ub') or '').lower().startswith(('year','yr')):value*=12
        if value>0:return value,f'read from {source}: \u201c{match.group().strip()[:60]}\u201d'
    return None,None

def normalized(fact,terms,item,rfx):
    """All arithmetic is deterministic. Unknown inputs produce no landed price."""
    out={'unit_price':None,'landed_unit_cost':None,'transformation_steps':[],'blocking_exception':False,'recurring':False}
    steps=out['transformation_steps']
    if fact.get('quote_status') in {'not_quoted','excluded'}:
        return out
    if fact.get('quote_status')=='unclear' or fact.get('confidence',0)<.85 or fact.get('mapping_issue'):
        return dict(out,blocking_exception=True,reason='Uncertain extraction or line mapping requires buyer review')
    if not fact.get('evidence_id'):
        return dict(out,blocking_exception=True,reason='Missing source evidence')
    raw=fact.get('raw_price'); basis=fact.get('raw_basis')
    currency=fact.get('raw_currency') or terms.get('currency') or rfx.get('currency')
    if currency=='mixed':
        return dict(out,blocking_exception=True,code='currency',reason='The response uses more than one currency and this line does not state which; confirm the quoting currency')
    if currency not in {'INR','USD'}:
        return dict(out,blocking_exception=True,code='currency',reason='Unsupported or missing currency')
    if not fact.get('raw_currency'):
        steps.append(f"Line does not state a currency; read as {currency} — {terms.get('currency_source','RFx quoting currency')}")
    if basis=='same_as_last_year':
        baseline=item.get('baseline_unit_price_inr')
        if not isinstance(baseline,(int,float)) or not math.isfinite(baseline) or baseline<=0:
            # Without a prior-year rate on file, "same as last year" is not a number.
            return dict(out,blocking_exception=True,code='baseline',reason='This event has no prior-year baseline for this line; enter the rate the supplier means as a correction')
        if not fact.get('prior_year_approved'):
            return dict(out,blocking_exception=True,code='baseline',reason='Confirm the buyer baseline is the intended prior-year rate')
        unit=baseline
        steps.append(f"Buyer-approved prior-year baseline: INR {unit}/piece; buyer reference")
    else:
        if not isinstance(raw,(int,float)) or not math.isfinite(raw) or raw<=0:
            return dict(out,blocking_exception=True,reason='Missing or invalid rate')
        if basis=='piece': unit=raw
        elif basis=='100_piece': unit=raw/100
        elif basis in {'carton','bundle'}:
            pack=fact.get('pack_size')
            if not isinstance(pack,(int,float)) or not math.isfinite(pack) or pack<=0:
                return dict(out,blocking_exception=True,reason='Missing or invalid pack size')
            unit=raw/pack
        elif basis=='kg':
            weight=item.get('unit_weight_kg') or fact.get('unit_weight_kg')
            if not isinstance(weight,(int,float)) or not math.isfinite(weight) or weight<=0:
                return dict(out,blocking_exception=True,code='weight',reason='Neither the requirement line nor the quote states a weight per unit, so a per-kilogram rate cannot be compared; enter the unit weight as a correction')
            unit=raw*weight
        elif basis=='other':
            covers=fact.get('covers_units')
            if not isinstance(covers,(int,float)) or isinstance(covers,bool) or not math.isfinite(covers) or covers<=0:
                return dict(out,blocking_exception=True,code='basis',reason='The quoted amount is stated as “'+str(fact.get('raw_basis_text') or 'an unstated basis')[:80]+'”; record how many '+str(item.get('unit') or 'units')+' one amount covers')
            unit=raw/covers
        else: return dict(out,blocking_exception=True,code='basis',reason='Unsupported or missing unit basis')
        steps.append(f"{raw} {currency} {fact.get('raw_basis_text') or basis}; " + (f"divide by {fact.get('pack_size')} pieces" if basis in {'carton','bundle'} else 'divide by 100' if basis=='100_piece' else f"multiply by {'buyer' if item.get('unit_weight_kg') else 'supplier-stated'} weight {weight} kg/piece" if basis=='kg' else f"divide by {fact.get('covers_units')} {item.get('unit') or 'units'} covered" if basis=='other' else f"per-{item.get('unit') or 'piece'} rate"))
        covers=fact.get('covers_units')
        if basis!='other' and isinstance(covers,(int,float)) and not isinstance(covers,bool) and math.isfinite(covers) and covers>1:
            unit/=covers
            steps.append(f"Amount covers {covers} {item.get('unit') or 'units'}; divide")
        period=fact.get('covers_period_months')
        if isinstance(period,(int,float)) and not isinstance(period,bool) and math.isfinite(period) and period>0:
            term,term_source=requirement_term(item,rfx)
            if not term:
                return dict(out,blocking_exception=True,code='term',reason='This is a recurring charge of '+str(fact.get('raw_basis_text') or f'{period:g} month(s)')[:70]+', but the requirement states no contract term; add the term in months to compare it')
            unit=unit/period*term
            out['recurring']=True
            steps.append(f"Recurring charge covering {period:g} month(s), extended over the {term:g}-month requirement term ({term_source})")
    if currency=='USD':
        fx=rfx['commercial_rules']['fx_rate_usd_inr']; unit*=fx
        steps.append(f'USD × {fx} INR/USD — fixed RFx comparison assumption, not a live market rate')
    out['unit_price']=round(unit,2)
    if terms.get('freight')=='included':
        landed=unit; steps.append('Freight included per supplier source')
    elif terms.get('freight')=='flat':
        landed=unit+terms['freight_per_piece_inr']; steps.append(f"Add freight INR {terms['freight_per_piece_inr']}/piece")
    elif terms.get('freight')=='percent':
        landed=unit*(1+terms['freight_percent']); steps.append(f"Add freight {terms['freight_percent']*100:g}% of material value")
    elif terms.get('freight')=='order_total':
        share=(terms.get('freight_shares') or {}).get(str(item['line_no']))
        if share is None:
            return dict(out,blocking_exception=True,code='freight',reason='Freight is quoted for the whole order but this line has no comparable value to apportion it against')
        landed=unit+share
        steps.append(f"Add freight {terms.get('freight_currency','INR')} {terms['freight_total']:,.2f} quoted for the whole order, apportioned across lines by value: INR {share:,.2f}/unit")
    elif out['recurring']:
        # A recurring charge is comparable on its own; any one-off freight is a term.
        landed=unit
        steps.append('Recurring charge compared excluding one-off freight; freight terms are listed beside the comparison')
    elif terms.get('freight')=='excluded':
        return dict(out,blocking_exception=True,code='freight',reason='Supplier states freight is charged separately without a rate; set the freight basis to compare landed cost')
    else:
        return dict(out,blocking_exception=True,code='freight',reason='Freight treatment is not stated; landed cost withheld')
    steps.append('GST excluded; conditional rebates excluded')
    out['landed_unit_cost']=round(landed,2)
    return out

# What a supporting document does not contain is not a question for the buyer when
# the supplier supplied it in the document that was meant to carry it.
DOCUMENT_GAP = re.compile(r'partial coverage|\b(?:unaddressed|unanswered|not answered)\b|'
                          r'\b(?:no|not|without)\b[^.;]{0,60}?\b(?:commercial\s+)?'
                          r'(?:quote|quotes|quotation|quoted|price|prices|pricing|rate|rates|line item|line items|'
                          r'answer|answers|answered|response|responses|addressed)\b',re.I)

def line_exposure(data):
    """Value at stake per line: the buyer baseline, or the quoted value when there is none."""
    per_piece={'piece':1,'100_piece':100,'carton':None,'bundle':None,'kg':None,'same_as_last_year':None}
    reference={}
    for item in data['rfx']['items']:
        unit=item.get('baseline_unit_price_inr') or 0
        if not unit:
            quotes=[f['raw_price']/per_piece[f.get('raw_basis')]
                    for v in data['vendors'].values() for f in v['facts']
                    if f.get('line_no')==item['line_no'] and per_piece.get(f.get('raw_basis'))
                    and isinstance(f.get('raw_price'),(int,float)) and f['raw_price']>0]
            unit=max(quotes) if quotes else 0
        reference[item['line_no']]=round(item['annual_quantity']*unit,2)
    return reference

def order_freight(data):
    """Spread a whole-order freight charge across the lines by their quoted value."""
    for vendor in data['vendors'].values():
        terms=vendor.get('terms') or {}
        terms.pop('freight_shares',None)
        if terms.get('freight')!='order_total' or not terms.get('freight_total'):continue
        neutral=dict(terms,freight='included')
        units={};total=0.0
        for item in data['rfx']['items']:
            fact=next((f for f in vendor['facts'] if f.get('line_no')==item['line_no']),None)
            if not fact:continue
            price=normalized(fact,neutral,item,data['rfx'])['unit_price']
            if price is None:continue
            units[item['line_no']]=price
            total+=price*(item.get('annual_quantity') or 0)
        if not total:continue
        amount=terms['freight_total']
        if terms.get('freight_currency')=='USD':amount*=data['rfx']['commercial_rules']['fx_rate_usd_inr']
        terms['freight_shares']={str(line):round(amount*price/total,4) for line,price in units.items()}

def rebuild(data):
    old={e['id']:e for e in data['exceptions']}
    exceptions=[]; rows=[]
    order_freight(data)
    reference=line_exposure(data)
    def issue(vid,line,kind,title,detail,eid=None,severity='high',blocking=False,confidence=.5):
        ident=f'{vid}:{line}:{kind}'
        previous=old.get(ident,{})
        exposure=reference.get(line,0) if line else (data['rfx']['expected_annual_spend_inr'] or round(sum(reference.values()),2))
        e={'id':ident,'vendor_id':vid,'line_no':line,'type':kind,'title':title,'detail':detail,'evidence_id':eid,'severity':severity,'financial_exposure_inr':exposure,'priority_score':round(exposure*(1-confidence)*(1 if data['vendors'][vid]['quality']=='pass' else .2),2),'status':previous.get('status','blocking' if blocking else 'needs_review')}
        if previous.get('resolution'):e['resolution']=previous['resolution']
        exceptions.append(e)
    for vid,v in data['vendors'].items():
        if v['response_status']=='awaiting_ingestion':continue
        if v['quality']!='pass':
            issue(vid,0,'qualification','Mandatory qualification '+v['quality'],
                  (v.get('quality_reason') or 'Supplier declarations require review').rstrip(' .')+'. Excluded from qualified-only awards.',
                  v.get('terms_evidence_id'),'critical',True)
        if v['terms'].get('rebate'):
            issue(vid,0,'conditional_discount','Conditional volume rebate',v['terms']['rebate']['source_text']+' — excluded from all award calculations.',v.get('terms_evidence_id'),'medium')
        # A certificate or authorisation letter was never going to carry prices. Saying so
        # is not a question for the buyer once the priced document has been read.
        supplied=any(f.get('quote_status')=='quoted' and f.get('raw_price') is not None for f in v['facts']) \
                 and bool(v.get('questionnaire'))
        notes=[n for n in (v.get('unresolved') or [])
               if not (supplied and DOCUMENT_GAP.search(n))]
        if notes:
            issue(vid,0,'document_review','Source interpretation notes','; '.join(notes),v.get('terms_evidence_id'),'medium')

    freight_blocked={}
    for item in data['rfx']['items']:
        row={'line_no':item['line_no'],'sku':item['sku'],'description':item['description'],'qty':item['annual_quantity'],'baseline_unit_price_inr':item['baseline_unit_price_inr'],'vendors':{}}
        for vid,v in data['vendors'].items():
            fact=next((f for f in v['facts'] if f['line_no']==item['line_no']),None)
            if not fact:
                status='awaiting_ingestion' if v['response_status']=='awaiting_ingestion' else 'unclear' if v.get('method') in {'vision_required','ai_required'} else 'not_quoted'
                q={'status':status,'qualification':v['quality'],'landed_unit_cost':None,'unit_price':None,'confidence':0,'raw_basis':'','evidence_id':v.get('terms_evidence_id'),'blocking_exception':status=='unclear'}
                if status in {'not_quoted','unclear'}:
                    issue(vid,item['line_no'],'missing_quote' if status=='not_quoted' else 'extraction_required',f"Line {item['line_no']}: {'not quoted' if status=='not_quoted' else 'extraction needed'}",'No usable price was extracted; never treated as zero.',v.get('terms_evidence_id'))
            else:
                n=normalized(fact,v['terms'],item,data['rfx'])
                q={**copy.deepcopy(fact),**n,'status':fact.get('quote_status','quoted'),'qualification':v['quality']}
                if n.get('code')=='freight':freight_blocked.setdefault(vid,[]).append(item['line_no'])
                if n['blocking_exception'] and n.get('code')!='freight':
                    issue(vid,item['line_no'],'review_price',f"Line {item['line_no']}: review price",n['reason'],fact['evidence_id'],'high',True,fact.get('confidence',.5))
                if fact.get('evidence_id'):
                    data['evidence'][fact['evidence_id']].update(line_no=item['line_no'],sku=item['sku'],description=item['description'],vendor_id=vid,vendor_name=v['name'],confidence=fact.get('confidence'),extraction_method=fact.get('extraction_method'),raw_fact=copy.deepcopy(fact),landed_unit_cost=n['landed_unit_cost'],normalized_unit_price=n['unit_price'],transformation_steps=n['transformation_steps'],buyer_override=fact.get('buyer_override'))
            row['vendors'][vid]=q
        rows.append(row)
    for vid,lines in freight_blocked.items():
        stated=data['vendors'][vid]['terms'].get('freight_source')
        issue(vid,0,'freight_basis','Freight basis needed for landed cost',
              ('Supplier states freight is charged separately: “'+str(stated)[:160]+'”. ' if stated else 'The response does not state whether freight is included. ')
              +f"Quoted prices are shown; landed cost on {len(lines)} line(s) stays withheld until you record the freight basis.",
              data['vendors'][vid].get('terms_evidence_id'),'high',True)
    # Retain resolved exception records even after the condition disappears.
    present={e['id'] for e in exceptions}
    exceptions.extend(e for k,e in old.items() if k not in present and e.get('status')=='resolved_by_buyer')
    data['comparison']=rows
    data['exceptions']=sorted(exceptions,key=lambda e:(e['status']=='resolved_by_buyer',-e['priority_score']))

def commit_extraction(vid,path,result,expected=None):
    digest=hashlib.sha256(path.read_bytes()).hexdigest()
    def mutate(data):
        v=data['vendors'][vid]
        # Name the document the supplier sent, not the file we stored it as.
        original=next((d.get('filename') for d in (v.get('documents') or []) if d.get('path')==str(path)),None) or path.name
        result['terms']=consolidate_currency(result.get('terms') or {},result.get('facts') or [],data['rfx'])
        # A genuinely new response invalidates prior review decisions for this supplier.
        # Their prior facts and decisions remain in immutable older snapshots.
        data['exceptions']=[e for e in data['exceptions'] if e['vendor_id']!=vid]
        v.update({k:copy.deepcopy(result[k]) for k in ['facts','terms','questionnaire','quality','method','unresolved','model']})
        v['quality_reason']=result.get('quality_reason') or ''
        v.update(file=original,format=path.suffix[1:],source_path=str(path),checksum=digest,response_status='processed',received_at=repo.now())
        v['lead']=next((f.get('lead_days') for f in v['facts'] if f.get('lead_days')),None)
        eid=str(uuid4());v['terms_evidence_id']=eid
        data['evidence'][eid]={'id':eid,'type':v['format'],'file':original,'source_path':str(path),'source_text':v['terms'].get('source_text','') or '; '.join(v['unresolved']),'vendor_id':vid,'vendor_name':v['name'],'blocks':result.get('blocks',[]),'extraction_method':v['method'],'checksum':digest}
        for fact in v['facts']:
            eid=str(uuid4());fact['evidence_id']=eid;fact['fact_id']=str(uuid4())
            data['evidence'][eid]={'id':eid,'type':v['format'],'file':path.name,'source_path':str(path),**fact.get('evidence',{}),'source_text':fact.get('evidence',{}).get('source_text') or fact.get('source_text',''),'checksum':digest}
        rebuild(data)
        data['stage']='exception_review'
    return repo.change('ingest:'+vid,'extractor',mutate,expected)

def ingest(vid,use_ai=True,path=None):
    data=repo.read()
    if data['rfx_status']!='approved':raise ValueError('Approve the RFx before collecting responses.')
    v=data['vendors'][vid]
    fixture=path is None
    path=Path(path) if path else DATA_DIR / repo.VENDORS[vid][1]
    digest=hashlib.sha256(path.read_bytes()).hexdigest()
    if v.get('checksum')==digest and v.get('method') not in {'vision_required','ai_required'}:
        return data  # preserve corrections on duplicate delivery/replayed graph node
    result=extract(vid,path,data['rfx'],use_ai,native_fixture=fixture)
    return commit_extraction(vid,path,result)

def resolve(eid,req):
    def mutate(data):
        exc=next((x for x in data['exceptions'] if x['id']==eid),None)
        if not exc:raise KeyError('Exception not found')
        if exc['status']=='resolved_by_buyer':raise ValueError('This exception is already resolved.')
        vid=exc['vendor_id'];v=data['vendors'][vid];line=exc['line_no'];action=req['action']
        reason=req.get('reason','').strip();actor=req.get('actor','').strip()
        if not reason or not actor:raise ValueError('Buyer name and reason are required.')
        fact=next((f for f in v['facts'] if f['line_no']==line),None)
        if line==0:
            if action=='freight':
                changes=req.get('changes',{});basis=changes.get('freight')
                if basis not in {'included','flat','percent'}:raise ValueError('Choose included, a flat rate per unit, or a percentage of material value.')
                amount=changes.get('value')
                terms=dict(v['terms'],freight=basis)
                terms.pop('freight_per_piece_inr',None);terms.pop('freight_percent',None)
                if basis in {'flat','percent'}:
                    if not isinstance(amount,(int,float)) or isinstance(amount,bool) or not math.isfinite(amount) or amount<=0:
                        raise ValueError('Enter the freight amount as a positive number.')
                    if basis=='flat':terms['freight_per_piece_inr']=float(amount)
                    else:
                        if amount>100:raise ValueError('Enter the freight percentage between 0 and 100.')
                        terms['freight_percent']=float(amount)/100
                terms['freight_source']=f'Buyer decision: {reason}'[:200]
                terms['freight_override']={'actor':actor,'reason':reason,'basis':basis,'value':amount,'timestamp':repo.now()}
                v['terms']=terms
            elif action=='qualify':
                if exc['type']!='qualification':raise ValueError('A qualification decision applies to the qualification review.')
                # The buyer is the authority on whether a declaration satisfies their requirement.
                v['quality']='pass'
                v['quality_reason']=f'Buyer confirmed qualification: {reason}'[:300]
                v['quality_override']={'actor':actor,'reason':reason,'timestamp':repo.now(),
                                       'previous_quality':exc['title'].split()[-1]}
            elif action!='accept':raise ValueError('Use acknowledge for supplier-level notes, or record a qualification decision.')
        else:
            if fact is None:
                fact={'line_no':line,'sku':data['rfx']['items'][line-1]['sku'],'raw_price':None,'raw_currency':'INR','raw_basis':'piece','pack_size':1,'lead_days':v['lead'],'confidence':0,'quote_status':'not_quoted','extraction_method':'buyer_entry'}
                v['facts'].append(fact)
            prior=copy.deepcopy(fact)
            if action=='accept':
                if fact.get('raw_basis')=='same_as_last_year':fact['prior_year_approved']=True
                elif fact.get('raw_price') is None:raise ValueError('An unreadable price cannot be accepted. Correct it with evidence or exclude it.')
                if fact.get('mapping_issue'):raise ValueError('Correct the line mapping before accepting this value.')
                fact.update(confidence=1,quote_status='quoted')
            elif action in {'correct','conversion'}:
                changes=req.get('changes',{})
                allowed={'raw_price','raw_currency','raw_basis','pack_size','unit_weight_kg','covers_units','covers_period_months','lead_days'}
                if not changes or set(changes)-allowed:raise ValueError('Provide supported raw-value corrections.')
                fact.update(changes);fact.update(confidence=1,quote_status='quoted')
            elif action=='remap':
                target=req.get('target_line')
                if not isinstance(target,int) or not 1<=target<=len(data['rfx']['items']):raise ValueError('Invalid target line')
                existing=next((f for f in v['facts'] if f['line_no']==target),None)
                if existing and existing.get('quote_status')!='excluded':raise ValueError('Target line already has a fact; exclude it before remapping.')
                if existing:v['facts'].remove(existing)
                fact.update(line_no=target,sku=data['rfx']['items'][target-1]['sku'])
                fact.pop('mapping_issue',None)
            elif action in {'not_quoted','exclude'}:
                fact['quote_status']='not_quoted' if action=='not_quoted' else 'excluded'
            else:raise ValueError('Unsupported resolution')
            new_eid=str(uuid4());fact['evidence_id']=new_eid;fact['fact_id']=str(uuid4())
            fact['buyer_override']={'actor':actor,'reason':reason,'action':action,'timestamp':repo.now(),'previous_fact':prior,'parent_evidence_id':prior.get('evidence_id')}
            base=copy.deepcopy(data['evidence'].get(prior.get('evidence_id'),{}))
            data['evidence'][new_eid]={**base,'id':new_eid,'source_text':base.get('source_text','') or reason,'buyer_override':fact['buyer_override']}
            if req.get('reply'):
                data['evidence'][new_eid]['clarification_reply']=req['reply']
            if action in {'correct','conversion','accept'}:
                norm=normalized(fact,v['terms'],data['rfx']['items'][fact['line_no']-1],data['rfx'])
                if norm['blocking_exception']:raise ValueError(norm.get('reason','Correction is still invalid'))
        exc['status']='resolved_by_buyer';exc['resolution']={**req,'at':repo.now()}
        rebuild(data)
    return repo.change('resolve:'+eid,req.get('actor','buyer'),mutate,req.get('expected_version'))
