import pytest
from fastapi.testclient import TestClient
from backend.app.main import app
from backend.app import repository as repo, intake, intake_docs, pipeline, graph, insights, jobs, knowledge
from backend.app.scenario import solve_scenario

@pytest.fixture
def client(tmp_path,monkeypatch):
    monkeypatch.setenv('AERCHAIN_DATA_DIR',str(tmp_path))
    monkeypatch.delenv('GOOGLE_API_KEY',raising=False)
    return TestClient(app)

def complete_intake(client,mode='custom',items=None):
    assert client.post('/api/intake/start',json={'mode':mode}).status_code==200
    fields={k:'Buyer-confirmed '+k for k in intake.FIELDS}
    fields.update(scope='Purchase ergonomic chairs',category='Office furniture',pricing='INR, freight included, GST excluded',compliance='No additional certificates required',deadline='2026-10-01 17:00 IST')
    if mode=='demo':fields.update(category='Corrugated packaging')
    body={'fields':fields,'confirm':True}
    if mode=='custom':body['items']=items or [{'description':'Ergonomic chair','quantity':120,'unit':'piece'}]
    assert client.post('/api/intake/save',json=body).status_code==200
    r=client.post('/api/intake/share',json={});assert r.status_code==200,r.text
    return r.json()

def test_checklist_gates_and_persistence(client):
    assert client.post('/api/intake/share',json={}).status_code==422
    assert client.post('/api/intake/save',json={'fields':{},'confirm':True}).status_code==422
    r=client.post('/api/intake/chat',json={'prompt':'I need office chairs'})
    assert r.status_code==200 and r.json()['fields']['scope']=='I need office chairs'
    assert len(client.get('/api/intake').json()['messages'])==2
    result=complete_intake(client)
    assert result['simulated'] and not result['suppliers']
    assert len(repo.read()['rfx']['items'])==1 and not repo.read()['vendors']
    assert client.post('/api/intake/chat',json={'prompt':'Change it'}).status_code==422
    assert client.post('/api/demo/inbox',json={}).status_code==422
    assert client.post('/api/intake/share',json={}).json()==result

def test_confirmation_invalidated_and_demo_mapping(client):
    client.post('/api/intake/start',json={'mode':'demo'})
    assert 'deadline' in client.get('/api/intake').json()['missing']
    client.post('/api/intake/save',json={'fields':{'deadline':'2026-10-01 5pm IST'},'confirm':True})
    assert client.get('/api/intake').json()['ready']
    client.post('/api/intake/save',json={'fields':{'budget':'Not disclosed'}})
    assert not client.get('/api/intake').json()['ready']
    assert client.post('/api/intake/share',json={}).status_code==422
    client.post('/api/intake/save',json={'fields':{},'confirm':True})
    assert len(client.post('/api/intake/share',json={}).json()['suppliers'])==5

def test_manual_supplier_dynamic_workflow(client,monkeypatch):
    complete_intake(client)
    r=client.post('/api/proposals',data={'supplier_name':'New Chairs Ltd','email':'quotes@example.com','body':'120 chairs; INR 2500 per chair, freight included; 10 days. Compliant.'})
    assert r.status_code==200,r.text
    vid=r.json()['vendor_id'];assert repo.read()['vendors'][vid]['email']=='quotes@example.com'
    assert not repo.read()['vendors'][vid]['facts']
    def extract(*args,**kwargs):
        return {'facts':[{'line_no':1,'sku':'ITEM-001','raw_price':2500,'raw_currency':'INR','raw_basis':'piece','pack_size':1,'lead_days':10,'quote_status':'quoted','confidence':.99,'source_text':'INR 2500 per chair'}],'terms':{'freight':'included','source_text':'freight included'},'questionnaire':{'Q1':{'answer':'Yes','confidence':.99}},'quality':'pass','method':'mock_extraction','unresolved':[],'blocks':[],'model':None}
    monkeypatch.setattr(pipeline,'extract',extract)
    graph.run('process')
    assert repo.read()['vendors'][vid]['response_status']=='processed'
    result=solve_scenario({},repo.read())
    assert result['award_total_inr']==300000 and result['savings_inr'] is None
    # New suppliers join an already checkpointed graph, not only its initial nodes.
    r=client.post('/api/proposals',data={'supplier_name':'Another Supplier','email':'other@example.com','body':'Proposal'})
    other=r.json()['vendor_id'];graph.run('process')
    assert repo.read()['vendors'][other]['response_status']=='processed'

def test_invalid_proposal_does_not_create_supplier(client):
    complete_intake(client);before=len(repo.read()['vendors'])
    for body in [{'supplier_name':'','email':'a@b.com','body':'x'},{'supplier_name':'Acme','email':'bad','body':'x'},{'supplier_name':'Acme','email':'a@b.com'}]:
        assert client.post('/api/proposals',data=body).status_code==422
    assert client.post('/api/proposals',data={'supplier_name':'Acme','email':'a@b.com'},files={'files':('bad.exe',b'bad')}).status_code==422
    assert len(repo.read()['vendors'])==before

def test_comparison_and_visual_routing(client):
    complete_intake(client,'demo');pipeline.receive_demo();graph.run('process')
    o=insights.overview()
    assert o['common_lines']==26 and o['leaders']['cost']
    assert all(x['past_deals'] is None for x in o['suppliers'])
    assert next(x for x in o['suppliers'] if x['id']=='boxworks')['cost'] is None
    for query,kind in [('Show a bar chart comparing cost','bar'),('Compare delivery lead time','bar'),('Compare quality','status'),('Compare previous deals','text'),('Give me a chart','text')]:
        r=client.post('/api/analysis',json={'question':query});assert r.status_code==200,r.text
        assert r.json()['kind']==kind
    r=client.post('/api/analysis',json={'question':'Show a pie chart of the cheapest qualified award'})
    assert r.status_code==200 and r.json()['chart']=='pie' and r.json()['scenario']['status']=='ok'

def test_intake_items_and_versions(client):
    client.post('/api/intake/start',json={'mode':'custom'})
    version=repo.read()['dataset_version']
    assert client.post('/api/intake/save',json={'fields':{},'items':[{'description':'item','quantity':0,'unit':'piece'}]}).status_code==422
    client.post('/api/intake/save',json={'fields':{'scope':'new'}})
    assert client.post('/api/intake/save',json={'fields':{'scope':'old'},'expected_version':version}).status_code==422

def requirement_sheet(path):
    from openpyxl import Workbook
    wb=Workbook();ws=wb.active
    for row in [['Office seating refresh'],['Category','Office furniture'],
                ['Delivery location','Bengaluru HQ, 4th floor'],['Payment terms','45 days from invoice'],
                ['Response deadline','15 November 2026, 17:00 IST'],['Offer validity','60 days'],
                ['Budget','INR 45 lakh indicative'],['Currency','INR, GST excluded, freight included'],
                ['Award criteria','Landed cost 60%, warranty 25%, lead time 15%'],
                ['Compliance','BIFMA or equivalent certificate mandatory'],
                ['Supplier preference','No preferred supplier'],['Contact','Priya Menon, priya@buyer.example'],
                ['Specifications','Adjustable lumbar, 5-year warranty'],[],
                ['SKU','Description','Annual Quantity','UOM'],
                ['CHR-ERG-001','Ergonomic task chair',120,'piece'],
                ['CHR-VIS-002','Visitor chair',60,'piece'],
                ['','',None,''],
                ['CHR-ERG-001','Ergonomic task chair',120,'piece']]:
        ws.append(row)
    wb.save(path);return path

def test_attachment_maps_sheet_onto_checklist(client,tmp_path):
    path=requirement_sheet(tmp_path/'requirements.xlsx')
    with open(path,'rb') as handle:
        r=client.post('/api/intake/attach',data={'prompt':'We need to refresh office seating at the Bengaluru HQ.'},
                      files={'files':('requirements.xlsx',handle,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert r.status_code==200,r.text
    value=r.json()
    # Every line is mapped once; the repeated row and the empty row are not requirement lines.
    # The buyer's own part number is kept as data, not glued into the description.
    assert [x['description'] for x in value['items']]==['Ergonomic task chair','Visitor chair']
    assert [x['sku'] for x in value['items']]==['CHR-ERG-001','CHR-VIS-002']
    assert value['items'][0]['quantity']==120 and value['items'][0]['unit']=='piece'
    assert value['fields']['payment']=='45 days from invoice'
    assert value['fields']['compliance']=='BIFMA or equivalent certificate mandatory'
    assert value['fields']['scope'].startswith('We need to refresh office seating')
    assert value['missing']==[] and value['confirmed'] is False
    assert value['messages'][-2]['attachments']==[{'filename':'requirements.xlsx','index':0}]
    assert client.get('/api/intake/documents/0').status_code==200
    assert client.get('/api/intake/documents/7').status_code==404
    # The header follows the buyer's own words, and confirmation is still explicit.
    # The header carries a short topic, not the buyer's whole sentence.
    title = client.get('/api/event').json()['title']
    assert 'office seating' in title.lower() and len(title.split()) <= 6
    assert not title.lower().startswith('we need')
    assert client.post('/api/intake/share',json={}).status_code==422
    assert client.post('/api/intake/save',json={'fields':{},'confirm':True}).status_code==200
    assert client.post('/api/intake/share',json={}).status_code==200

def test_first_open_starts_empty(client):
    d=client.get('/api/state').json()
    assert d['rfx']['title']=='' and d['rfx']['items']==[] and d['vendors']=={}
    assert client.get('/api/intake').json()['messages']==[]
    assert client.post('/api/demo/inbox',json={}).status_code==422
    # The bundled packaging event is opt-in, and its suppliers appear only once their
    # demo messages are loaded: an inbox never shows a supplier who has sent nothing.
    assert client.post('/api/intake/start',json={'mode':'demo'}).status_code==200
    d=client.get('/api/state').json()
    assert len(d['rfx']['items'])==30 and d['vendors']=={}
    complete_intake(client,'demo')
    assert client.get('/api/state').json()['vendors']=={}
    assert client.post('/api/demo/inbox',json={}).status_code==200
    assert len(client.get('/api/state').json()['vendors'])==5

def test_attachment_rejects_unreadable_and_oversized_packets(client,tmp_path):
    path=tmp_path/'notes.rtf';path.write_text('Category: Office furniture')
    with open(path,'rb') as handle:
        r=client.post('/api/intake/attach',data={'prompt':'x'},files={'files':('notes.rtf',handle,'application/rtf')})
    assert r.status_code==422 and 'Excel' in r.json()['detail']
    assert client.post('/api/intake/attach',data={'prompt':'no files'}).status_code==422
    assert client.get('/api/intake').json()['messages']==[]

FREIGHT_PHRASES = [
    ('Prices are inclusive of delivery to site', 'included'),
    ('All rates include freight to Bengaluru', 'included'),
    ('Freight included', 'included'),
    ('DDP Bengaluru plant', 'included'),
    ('F.O.R. destination', 'included'),
    ('Free delivery within city limits', 'included'),
    ('Ex-works Pune', 'excluded'),
    ('Prices are FOB Chennai', 'excluded'),
    ('Freight extra at actuals', 'excluded'),
    ('Delivery charges additional', 'excluded'),
    ('Freight will be borne by the buyer', 'excluded'),
    ('Our quote covers the goods only', 'unknown'),
    ('Lead time 21 days', 'unknown'),
]

@pytest.mark.parametrize('text,expected',FREIGHT_PHRASES)
def test_freight_wording_is_recognised(text,expected):
    # A landed cost was withheld for every supplier who did not use the words "freight included".
    assert pipeline.terms_from_text(text)['freight']==expected

def quote_extraction(freight_text,price,answers):
    def extract(*args,**kwargs):
        return {'facts':[{'line_no':1,'sku':'ITEM-001','raw_price':price,'raw_currency':'INR','raw_basis':'piece',
                          'pack_size':1,'lead_days':12,'quote_status':'quoted','confidence':.99,'source_text':f'INR {price} per chair'}],
                'terms':pipeline.terms_from_text(freight_text),
                'questionnaire':{k:{'answer':v,'confidence':.99} for k,v in answers.items()},
                'quality':'pass','method':'mock_extraction','unresolved':[],'blocks':[],'model':None,
                'captured_terms':[{'term_type':'payment_terms','value':'30 days from invoice','confidence':.9,'source_text':'Payment: 30 days'},
                                  {'term_type':'warranty','value':'3 years on frames','confidence':.9,'source_text':'Warranty: 3 years'}]}
    return extract

def test_stated_freight_exclusion_is_a_buyer_decision(client,monkeypatch):
    complete_intake(client)
    vid=client.post('/api/proposals',data={'supplier_name':'SeatSmart','email':'bids@example.com',
                    'body':'Ex-works Hosur; freight extra at actuals. INR 7990 per chair.'}).json()['vendor_id']
    monkeypatch.setattr(pipeline,'extract',quote_extraction('Ex-works Hosur; freight extra at actuals',7990,{'Q1':'Yes'}))
    graph.run('process')
    d=repo.read()
    row=d['comparison'][0]['vendors'][vid]
    assert row['unit_price']==7990 and row['landed_unit_cost'] is None
    # One supplier-level decision, not the same note repeated on every line.
    freight=[e for e in d['exceptions'] if e['type']=='freight_basis']
    assert len(freight)==1 and freight[0]['line_no']==0 and freight[0]['status']=='blocking'
    assert not [e for e in d['exceptions'] if e['type']=='review_price']
    assert freight[0]['financial_exposure_inr']>0, 'an event without a baseline still needs a priority'
    r=client.post('/api/exceptions/'+freight[0]['id']+'/resolve',json={'action':'freight','actor':'Priya Menon',
        'reason':'Supplier confirmed road freight is 3% of material value.','expected_version':repo.read()['dataset_version'],
        'changes':{'freight':'percent','value':3}})
    assert r.status_code==200,r.text
    d=repo.read()
    assert d['comparison'][0]['vendors'][vid]['landed_unit_cost']==round(7990*1.03,2)
    assert d['vendors'][vid]['terms']['freight_override']['actor']=='Priya Menon'
    assert not [e for e in d['exceptions'] if e['type']=='freight_basis' and e['status']!='resolved_by_buyer']

def test_freight_decision_validates_its_input(client,monkeypatch):
    complete_intake(client)
    client.post('/api/proposals',data={'supplier_name':'SeatSmart','email':'bids@example.com','body':'Ex-works Hosur.'})
    monkeypatch.setattr(pipeline,'extract',quote_extraction('Ex-works Hosur',7990,{'Q1':'Yes'}))
    graph.run('process')
    eid=next(e['id'] for e in repo.read()['exceptions'] if e['type']=='freight_basis')
    body=lambda changes:{'action':'freight','actor':'Priya Menon','reason':'Checked with the supplier.',
                         'expected_version':repo.read()['dataset_version'],'changes':changes}
    for changes in [{'freight':'percent','value':0},{'freight':'percent','value':180},{'freight':'flat','value':-2},
                    {'freight':'flat'},{'freight':'guess','value':3}]:
        assert client.post('/api/exceptions/'+eid+'/resolve',json=body(changes)).status_code==422,changes
    assert repo.read()['comparison'][0]['vendors'][next(iter(repo.read()['vendors']))]['landed_unit_cost'] is None

def test_the_questions_asked_are_captured_and_compared(client,monkeypatch):
    complete_intake(client)
    rfx=client.get('/api/rfx').json()
    asked={q['id']:q['question'].lower() for q in rfx['questionnaire']}
    # What the buyer stated on the checklist is what suppliers are asked about.
    assert len(asked)>=8 and sum(q['mandatory'] for q in rfx['questionnaire'])==1
    assert any('freight' in q for q in asked.values()) and any('payment' in q for q in asked.values())
    assert any('warranty' in q for q in asked.values()) and any('valid' in q for q in asked.values())
    client.post('/api/proposals',data={'supplier_name':'ErgoWorks','email':'sales@example.com','body':'Prices include delivery.'})
    monkeypatch.setattr(pipeline,'extract',quote_extraction('Prices are inclusive of delivery to site',8450,
                        {'Q1':'Yes','Q3':'Yes, delivery included','Q5':'45 days from invoice','Q6':'60','Q8':'5 years'}))
    graph.run('process')
    o=client.get('/api/overview').json()
    rows={r['key'] for r in o['term_rows']}
    assert {'freight','payment','warranty'} <= rows
    # Offer validity and stated lead time repeat the table, so the comparison drops them...
    assert 'validity' not in rows and 'lead_time' not in rows
    terms=o['suppliers'][0]['terms']
    assert terms['freight']['value']=='Included in price'
    assert terms['payment']['value']=='45 days from invoice' and terms['warranty']['value']=='5 years'
    # ...while staying answerable in the analysis room.
    assert terms['validity']['value']=='60'
    assert 'term:validity' in {d['key'] for d in client.get('/api/graph/schema').json()['dimensions']}
    assert len(o['questions'])==len(rfx['questionnaire'])
    assert o['suppliers'][0]['cost']==8450*120
    # Those same facts answer a question in the analysis room.
    answer=client.post('/api/analysis',json={'question':'Compare supplier payment terms'}).json()
    assert answer['kind']=='facts' and answer['points'][0]['value']=='45 days from invoice'
    assert client.post('/api/analysis',json={'question':'Compare warranty across suppliers'}).json()['metric']=='warranty'

def priced(line_no=1,**fact):
    base={'line_no':line_no,'sku':f'ITEM-{line_no:03d}','raw_price':40.0,'raw_currency':None,'raw_basis':'kg',
          'pack_size':None,'unit_weight_kg':None,'lead_days':12,'quote_status':'quoted','confidence':.99,
          'source_text':'CB-5P-001: 40.00/kg'}
    return {**base,**fact}

def extraction(facts,text='Freight included.'):
    def extract(*args,**kwargs):
        return {'facts':[dict(f) for f in facts],'terms':pipeline.terms_from_text(text),
                'questionnaire':{'Q1':{'answer':'Yes','confidence':.99}},'quality':'pass','method':'mock',
                'unresolved':[],'blocks':[],'model':None,'captured_terms':[]}
    return extract

def one_supplier(client,monkeypatch,facts,text='Freight included.',items=None):
    complete_intake(client,items=items)
    vid=client.post('/api/proposals',data={'supplier_name':'Boxwork','email':'sales@example.com','body':text}).json()['vendor_id']
    monkeypatch.setattr(pipeline,'extract',extraction(facts,text))
    graph.run('process')
    return vid

def cell(vid,line_no=1):
    row=next(r for r in repo.read()['comparison'] if r['line_no']==line_no)
    return row['vendors'][vid]

def test_currency_stated_once_applies_to_every_line(client,monkeypatch):
    # A quote that names INR in its header does not repeat it on each row.
    vid=one_supplier(client,monkeypatch,[priced(raw_basis='piece',raw_price=14.05)],
                     'All rates in INR. Freight included.')
    q=cell(vid)
    assert q['landed_unit_cost']==14.05
    assert any('does not state a currency' in step for step in q['transformation_steps'])
    assert repo.read()['vendors'][vid]['terms']['currency']=='INR'

def test_currency_falls_back_to_the_rfx_when_nobody_states_one(client,monkeypatch):
    vid=one_supplier(client,monkeypatch,[priced(raw_basis='piece',raw_price=14.05)],'Freight included.')
    assert cell(vid)['landed_unit_cost']==14.05
    assert repo.read()['vendors'][vid]['terms']['currency_source'].startswith('not stated')

def test_a_dollar_quote_is_never_read_as_rupees(client,monkeypatch):
    vid=one_supplier(client,monkeypatch,[priced(raw_basis='piece',raw_price=2.0)],
                     'All prices in USD. Freight included.')
    q=cell(vid)
    assert repo.read()['vendors'][vid]['terms']['currency']=='USD'
    assert q['landed_unit_cost']==round(2.0*repo.read()['rfx']['commercial_rules']['fx_rate_usd_inr'],2)

def test_a_mixed_currency_response_is_not_guessed(client,monkeypatch):
    vid=one_supplier(client,monkeypatch,[priced(raw_basis='piece',raw_price=14.05)],
                     'Domestic lines in INR, imported lines in USD. Freight included.')
    q=cell(vid)
    assert q['landed_unit_cost'] is None and 'more than one currency' in q['reason']

def test_per_kilogram_rate_uses_a_supplier_or_buyer_weight(client,monkeypatch):
    # The supplier states the weight themselves.
    vid=one_supplier(client,monkeypatch,[priced(unit_weight_kg=0.62)])
    assert cell(vid)['landed_unit_cost']==round(40*0.62,2)
    assert any('supplier-stated weight' in s for s in cell(vid)['transformation_steps'])

def test_per_kilogram_rate_uses_the_weight_on_the_requirement_line(client,monkeypatch):
    vid=one_supplier(client,monkeypatch,[priced()],
                     items=[{'description':'5-ply RSC carton','quantity':120,'unit':'piece','unit_weight_kg':0.62}])
    assert cell(vid)['landed_unit_cost']==round(40*0.62,2)
    assert any('buyer weight' in s for s in cell(vid)['transformation_steps'])

def test_a_weightless_kilogram_rate_is_withheld_until_the_buyer_supplies_one(client,monkeypatch):
    vid=one_supplier(client,monkeypatch,[priced()])
    q=cell(vid)
    assert q['unit_price'] is None and 'weight per unit' in q['reason']
    eid=next(e['id'] for e in repo.read()['exceptions'] if e['line_no']==1 and e['type']=='review_price')
    r=client.post('/api/exceptions/'+eid+'/resolve',json={'action':'conversion','actor':'Priya Menon',
        'reason':'Buyer spec sheet states the finished box weight.','expected_version':repo.read()['dataset_version'],
        'changes':{'raw_price':40.0,'raw_currency':'INR','raw_basis':'kg','pack_size':1,'unit_weight_kg':0.62}})
    assert r.status_code==200,r.text
    assert cell(vid)['landed_unit_cost']==round(40*0.62,2)

def test_same_as_last_year_never_becomes_zero(client,monkeypatch):
    # This event has no prior-year rate, so "same as last year" is not a number.
    vid=one_supplier(client,monkeypatch,[priced(raw_basis='same_as_last_year',raw_price=None)])
    q=cell(vid)
    assert q['landed_unit_cost'] is None and 'no prior-year baseline' in q['reason']
    eid=next(e['id'] for e in repo.read()['exceptions'] if e['line_no']==1 and e['type']=='review_price')
    r=client.post('/api/exceptions/'+eid+'/resolve',json={'action':'accept','actor':'Priya Menon',
        'reason':'Treat as the prior-year rate.','expected_version':repo.read()['dataset_version']})
    assert r.status_code==422 and 'no prior-year baseline' in r.json()['detail']
    assert cell(vid)['landed_unit_cost'] is None

def test_a_pending_supplier_is_still_priced_but_never_leads(client,monkeypatch):
    complete_intake(client)
    slow=client.post('/api/proposals',data={'supplier_name':'Packright','email':'a@example.com','body':'Freight included.'}).json()['vendor_id']
    fast=client.post('/api/proposals',data={'supplier_name':'Alpha','email':'b@example.com','body':'Freight included.'}).json()['vendor_id']
    def extract(vendor_id,*args,**kwargs):
        cheap=vendor_id==slow
        return {'facts':[priced(raw_basis='piece',raw_price=39.5 if cheap else 44.0,raw_currency='INR')],
                'terms':pipeline.terms_from_text('Freight included.'),
                'questionnaire':{} if cheap else {'Q1':{'answer':'Yes','confidence':.99}},
                'quality':'pending' if cheap else 'pass','method':'mock','unresolved':[],'blocks':[],'model':None,'captured_terms':[]}
    monkeypatch.setattr(pipeline,'extract',extract)
    graph.run('process')
    o=insights.overview()
    costs={v['id']:v['cost'] for v in o['suppliers']}
    # The cheaper supplier is pending, so its price is shown but the green marker is not.
    assert costs[slow] is not None and costs[slow]<costs[fast]
    assert o['leaders']['cost']==[fast]
    award=solve_scenario({},repo.read())
    assert [m['vendor_id'] for m in award['vendor_mix']]==[fast]

def test_a_weight_column_on_the_requirement_sheet_is_mapped(client,tmp_path):
    from openpyxl import Workbook
    wb=Workbook();ws=wb.active
    for row in [['SKU','Description','Annual Quantity','UOM','Unit weight (kg)'],
                ['CB-5P-001','5-ply RSC carton',120000,'piece',0.62],
                ['CB-3P-006','3-ply RSC carton',210000,'piece','0.21']]:
        ws.append(row)
    path=tmp_path/'lines.xlsx';wb.save(path)
    with open(path,'rb') as handle:
        r=client.post('/api/intake/attach',data={'prompt':'Corrugated boxes for the Pune plant.'},
                      files={'files':('lines.xlsx',handle)})
    assert r.status_code==200,r.text
    assert [x.get('unit_weight_kg') for x in r.json()['items']]==[0.62,0.21]

RENTAL_LINE=[{'description':'Apple MacBook Pro 16-inch (60-month rental)','quantity':100,'unit':'device'}]

def rental(**fact):
    base={'line_no':1,'sku':'ITEM-001','raw_price':8950,'raw_currency':'INR','raw_basis':'piece','pack_size':None,
          'unit_weight_kg':None,'raw_basis_text':'per device per month','covers_units':1,'covers_period_months':1,
          'lead_days':21,'quote_status':'quoted','confidence':.99,'source_text':'Monthly rental / device (INR) | 8950'}
    return {**base,**fact}

def test_a_certificate_attached_beside_a_proposal_keeps_the_price(client,monkeypatch):
    # The ISO certificate is read first and has no pricing; it must not overwrite the quote.
    complete_intake(client,items=RENTAL_LINE)
    vid=client.post('/api/proposals',data={'supplier_name':'NorthStar','email':'a@example.com',
        'body':'Proposal attached. Freight included.'}).json()['vendor_id']
    documents=[]
    def extract(vendor_id,path,*args,**kwargs):
        documents.append(path.name)
        certificate=len(documents)==1
        return {'facts':[rental(quote_status='unclear',raw_price=None,raw_currency=None,raw_basis=None,
                                covers_units=None,covers_period_months=None,confidence=.95)] if certificate else [rental()],
                'terms':pipeline.terms_from_text('Freight included.'),
                'questionnaire':{} if certificate else {'Q1':{'answer':'Yes','confidence':.99}},
                'quality':'pass','method':'mock','unresolved':['No pricing in this document'] if certificate else [],
                'blocks':[],'model':None,'captured_terms':[]}
    monkeypatch.setattr(pipeline,'extract',extract)
    # Two attachments: a certificate, then the commercial proposal.
    import backend.app.pipeline as pl
    repo.change('two_documents','test',lambda d:d['vendors'][vid].update(
        documents=[{'path':d['vendors'][vid]['documents'][0]['path'],'filename':'ISO9001.pdf'},
                   {'path':d['vendors'][vid]['documents'][0]['path'],'filename':'Proposal.xlsx'}]))
    pl.ingest_received(vid,use_ai=False)
    fact=repo.read()['vendors'][vid]['facts'][0]
    assert fact['raw_price']==8950 and fact['quote_status']=='quoted' and not fact.get('mapping_issue')
    # And the certificate's own note says which document it came from.
    assert any(n.startswith('ISO9001.pdf:') for n in repo.read()['vendors'][vid]['unresolved'])

def test_a_monthly_rental_is_compared_over_the_requirement_term(client,monkeypatch):
    vid=one_supplier(client,monkeypatch,[rental()],items=RENTAL_LINE)
    q=cell(vid)
    # 8,950 per device per month across the 60 months the requirement states.
    assert q['landed_unit_cost']==8950*60
    assert any('60-month requirement term' in s for s in q['transformation_steps'])
    assert any('read from the requirement line' in s for s in q['transformation_steps'])

def test_a_fleet_price_per_quarter_reaches_the_buyer_unit(client,monkeypatch):
    vid=one_supplier(client,monkeypatch,[rental(raw_price=2745000,raw_basis='other',
        raw_basis_text='per quarter for all 100 devices',covers_units=100,covers_period_months=3)],items=RENTAL_LINE)
    q=cell(vid)
    assert q['landed_unit_cost']==round(2745000/100/3*60,2)
    assert any('divide by 100 device covered' in s for s in q['transformation_steps'])

def test_a_dollar_rental_converts_once(client,monkeypatch):
    vid=one_supplier(client,monkeypatch,[rental(raw_price=103.0,raw_currency='USD')],items=RENTAL_LINE)
    fx=repo.read()['rfx']['commercial_rules']['fx_rate_usd_inr']
    assert cell(vid)['landed_unit_cost']==round(103.0*fx*60,2)

def test_a_recurring_charge_is_not_withheld_for_unstated_freight(client,monkeypatch):
    # A monthly rental fee is not a delivered-goods price.
    vid=one_supplier(client,monkeypatch,[rental()],text='Rental quote attached.',items=RENTAL_LINE)
    assert repo.read()['vendors'][vid]['terms']['freight']=='unknown'
    assert cell(vid)['landed_unit_cost']==8950*60
    assert not [e for e in repo.read()['exceptions'] if e['type']=='freight_basis']
    # A one-off goods price still waits for the freight basis.
    assert cell(vid)['recurring'] is True

def test_a_one_off_goods_price_still_waits_for_freight(client,monkeypatch):
    vid=one_supplier(client,monkeypatch,[rental(raw_basis='piece',raw_price=14.05,raw_basis_text='per piece',
        covers_units=1,covers_period_months=None)],text='Quote attached.')
    assert cell(vid)['landed_unit_cost'] is None
    assert [e for e in repo.read()['exceptions'] if e['type']=='freight_basis']

def test_a_recurring_charge_without_a_term_asks_for_one(client,monkeypatch):
    vid=one_supplier(client,monkeypatch,[rental()],
                     items=[{'description':'Laptop supply','quantity':100,'unit':'device'}])
    q=cell(vid)
    assert q['landed_unit_cost'] is None and 'states no contract term' in q['reason']
    eid=next(e['id'] for e in repo.read()['exceptions'] if e['type']=='review_price')
    assert client.post('/api/exceptions/'+eid+'/resolve',json={'action':'conversion','actor':'Priya Menon',
        'reason':'Supplier confirmed the rate covers the full 36-month term.','expected_version':repo.read()['dataset_version'],
        'changes':{'raw_price':8950,'raw_currency':'INR','raw_basis':'piece','pack_size':1,
                   'covers_units':1,'covers_period_months':None}}).status_code==200
    assert cell(vid)['landed_unit_cost']==8950

def test_an_unexplained_basis_quotes_the_supplier_back(client,monkeypatch):
    vid=one_supplier(client,monkeypatch,[rental(raw_basis='other',raw_basis_text='per pallet',
        covers_units=None,covers_period_months=None)],items=RENTAL_LINE)
    q=cell(vid)
    assert q['landed_unit_cost'] is None
    assert 'per pallet' in q['reason'] and 'how many device' in q['reason']

def test_a_term_in_the_requirement_text_is_read_but_a_warranty_is_not():
    term=pipeline.requirement_term
    assert term({'description':'MacBook Pro (60-month rental)'},{})[0]==60
    assert term({'description':'Laptop'},{'scope':'Five-year rental over a 60-month rental term.'})[0]==60
    assert term({'description':'Managed service, 3 year contract'},{})[0]==36
    assert term({'description':'Chair with 12-month reseller warranty'},{'scope':'Office chairs'})[0] is None
    assert term({'description':'Corrugated box 450x300'},{'scope':'Annual packaging'})[0] is None
    assert term({'description':'x','term_months':24},{})==(24.0,'stated on the requirement line')

def test_a_term_column_on_the_requirement_sheet_is_mapped(client,tmp_path):
    from openpyxl import Workbook
    wb=Workbook();ws=wb.active
    for row in [['Description','Quantity','UOM','Rental term (months)'],['MacBook Pro 16-inch',100,'device',60]]:
        ws.append(row)
    path=tmp_path/'rental.xlsx';wb.save(path)
    with open(path,'rb') as handle:
        r=client.post('/api/intake/attach',data={'prompt':'Five-year laptop rental.'},files={'files':('rental.xlsx',handle)})
    assert r.status_code==200,r.text
    assert r.json()['items'][0]['term_months']==60

def two_suppliers(client,monkeypatch):
    """One cheaper, one faster, both qualified — the shape of a real trade-off."""
    complete_intake(client,items=RENTAL_LINE)
    cheap=client.post('/api/proposals',data={'supplier_name':'NorthStar','email':'a@example.com','body':'Freight included.'}).json()['vendor_id']
    fast=client.post('/api/proposals',data={'supplier_name':'QuantumEdge','email':'b@example.com','body':'Freight included.'}).json()['vendor_id']
    def extract(vendor_id,*args,**kwargs):
        is_cheap=vendor_id==cheap
        return {'facts':[rental(raw_price=8950 if is_cheap else 9150,lead_days=21 if is_cheap else 14)],
                'terms':pipeline.terms_from_text('Freight included.'),
                'questionnaire':{'Q1':{'answer':'Yes','confidence':.99},
                                 'Q5':{'answer':'Monthly, Net 30' if is_cheap else 'Quarterly in advance','confidence':.9}},
                'quality':'pass','method':'mock','unresolved':[],'blocks':[],'model':None,
                'captured_terms':[{'term_type':'warranty','value':'3 years' if is_cheap else '5 years','confidence':.9,'source_text':'w'}]}
    monkeypatch.setattr(pipeline,'extract',extract)
    graph.run('process')
    return cheap,fast

def test_the_graph_is_built_from_the_snapshot(client,monkeypatch):
    cheap,fast=two_suppliers(client,monkeypatch)
    g=knowledge.build()
    counts=g.as_dict()['counts']
    assert counts['event']==1 and counts['supplier']==2 and counts['requirement']==1
    assert counts['question']==len(repo.read()['rfx']['questionnaire'])
    assert counts['quote']==2 and counts['term']>=2 and counts['answer']>=2
    # Every quote hangs off its supplier and points at the requirement it prices.
    quote=g.out(f'supplier:{cheap}','quoted')[0]
    assert g.out(quote['id'],'for_requirement')[0]['id']=='requirement:1'
    assert g.out(quote['id'],'evidenced_by'), 'a quote must keep its source'
    assert quote['attrs']['landed_unit_cost']==8950*60
    assert g.out(f'supplier:{cheap}','responds_to')[0]['id']=='event'
    # An answer is linked to the question it answers, not just stored beside it.
    answer=next(a for a in g.out(f'supplier:{cheap}','answered') if a['attrs']['question_id']=='Q1')
    assert g.out(answer['id'],'answers')[0]['id']=='question:Q1'

def test_the_graph_versions_with_the_data(client,monkeypatch):
    two_suppliers(client,monkeypatch)
    current=repo.read()['dataset_version']
    assert knowledge.build().version==current
    # An older snapshot still builds its own graph, with no suppliers yet.
    early=knowledge.build(repo.read(1))
    assert early.version==1 and not early.of_type('supplier')

def test_dimensions_come_from_the_event_not_a_fixed_list(client,monkeypatch):
    two_suppliers(client,monkeypatch)
    keys={d['key'] for d in knowledge.dimensions(knowledge.build())}
    assert {'cost','delivery','coverage','qualification','reviews'} <= keys
    assert 'term:warranty' in keys and 'term:payment' in keys
    # Every question the RFx asked is comparable in its own right.
    for question in repo.read()['rfx']['questionnaire']:
        assert f"answer:{question['id']}" in keys

def test_a_query_compares_only_the_suppliers_asked_about(client,monkeypatch):
    cheap,fast=two_suppliers(client,monkeypatch)
    third=client.post('/api/proposals',data={'supplier_name':'Ignored Ltd','email':'c@example.com','body':'x'}).json()['vendor_id']
    result=knowledge.compare({'suppliers':[f'supplier:{cheap}',f'supplier:{fast}'],'dimensions':['cost','delivery']})
    assert [s['id'] for s in result['suppliers']]==[f'supplier:{cheap}',f'supplier:{fast}']
    assert third not in str(result['suppliers'])
    cost,delivery=result['rows']
    assert cost['winners']==[f'supplier:{cheap}'] and delivery['winners']==[f'supplier:{fast}']
    assert cost['cells'][f'supplier:{cheap}']['numeric']==8950*60*100
    # The verdict only restates what the table computed, in one short line.
    text=knowledge.verdict(result)
    assert 'NorthStar: cost' in text and 'QuantumEdge: lead time' in text
    assert len(text.split())<=30, f'verdict must stay short, got {len(text.split())} words: {text}'

def test_a_query_defaults_to_a_rounded_comparison(client,monkeypatch):
    two_suppliers(client,monkeypatch)
    result=knowledge.compare({})
    keys=[r['key'] for r in result['rows']]
    assert 'cost' in keys and 'qualification' in keys and any(k.startswith('term:') for k in keys)
    assert len(result['suppliers'])==2

def test_the_comparable_basket_is_the_lines_both_suppliers_priced(client,monkeypatch):
    cheap,fast=two_suppliers(client,monkeypatch)
    g=knowledge.build()
    assert knowledge.basket(g,[f'supplier:{cheap}',f'supplier:{fast}'])==[1]
    # Withdraw one supplier's usable price and the shared basket empties honestly.
    repo.change('blank','test',lambda d:d['vendors'][fast]['facts'].clear() or pipeline.rebuild(d))
    result=knowledge.compare({'suppliers':[f'supplier:{cheap}',f'supplier:{fast}'],'dimensions':['cost']})
    assert result['compared_lines']==[]
    assert result['rows'][0]['cells'][f'supplier:{cheap}']['display']=='Not comparable'
    text=knowledge.verdict(result)
    assert 'no line priced by all' in text and len(text.split())<=30

def test_the_graph_endpoints_serve_the_same_structure(client,monkeypatch):
    cheap,_=two_suppliers(client,monkeypatch)
    whole=client.get('/api/graph').json()
    assert whole['counts']['supplier']==2 and whole['edges']
    schema=client.get('/api/graph/schema').json()
    assert {s['name'] for s in schema['suppliers']}=={'NorthStar','QuantumEdge'}
    assert any(d['key']=='term:warranty' for d in schema['dimensions'])
    node=client.get(f'/api/graph/nodes/supplier:{cheap}').json()
    assert node['type']=='supplier' and node['edges']
    assert client.get('/api/graph/nodes/supplier:nobody').status_code==404
    answer=client.post('/api/graph/query',json={'suppliers':[f'supplier:{cheap}'],'dimensions':['cost','term:warranty']}).json()
    assert [r['key'] for r in answer['rows']]==['cost','term:warranty']
    assert answer['rows'][1]['cells'][f'supplier:{cheap}']['display']=='3 years'
    assert answer['verdict']


def test_the_assistant_keeps_its_replies_short(client,tmp_path):
    """Two complaints about verbosity is two too many; the cap is now a test."""
    from openpyxl import Workbook
    wb=Workbook();ws=wb.active
    for row in [['Category','IT hardware rental'],['Delivery location','Bengaluru office'],
                ['Payment terms','Monthly in arrears'],['Currency','INR, GST excluded'],
                ['Compliance','ISO 9001'],['Specifications','16-inch, M4'],[],
                ['SKU','Description','Quantity','UOM'],['MBP-16-M4','MacBook Pro 16-inch',100,'device']]:
        ws.append(row)
    path=tmp_path/'req.xlsx';wb.save(path)
    with open(path,'rb') as handle:
        value=client.post('/api/intake/attach',data={'prompt':'Requirement sheet attached.'},
                          files={'files':('01_Buyer_RFx.xlsx',handle)}).json()
    reply=value['messages'][-1]['text']
    assert len(reply.split())<=intake.WORD_LIMIT, f'{len(reply.split())} words: {reply}'
    # It names what is still missing and asks for it in plain words.
    assert value['missing']==['budget','deadline']
    assert 'budget' in reply.lower() and 'deadline' in reply.lower()
    assert 'checklist field' not in reply.lower() and 'mapped from' not in reply.lower()

def test_the_opening_line_is_a_question_not_a_briefing(client):
    client.post('/api/intake/start',json={'mode':'custom'})
    greeting=client.get('/api/intake').json()['messages'][0]['text']
    assert len(greeting.split())<=25 and greeting.endswith('.')
    assert '?' in greeting

def test_a_long_model_reply_is_trimmed_to_whole_sentences():
    long_reply=('Data mapped from the buyer RFx and the supplier response template: scope, category, '
                'specifications, delivery, pricing, payment, compliance and item line details. '
                'The following items remain missing: budget and response deadline. '
                'Please review the mapped details and provide the missing items so we can finalise '
                'requirements for supplier matching and proceed to the next stage of this event.')
    trimmed=intake.brief(long_reply,limit=30)
    assert len(trimmed.split())<=30 and trimmed.endswith('.')
    assert trimmed in long_reply, 'trimming must not invent words'
    assert intake.brief('Short and sweet.')=='Short and sweet.'

def test_every_checklist_item_carries_a_plain_question(client):
    for item in client.get('/api/intake').json()['checklist']:
        assert item['ask'] and len(item['ask'].split())<=16, item
        assert item['ask'].endswith(('?','.')), item


def test_a_blank_supplier_template_does_not_win_over_the_real_table(client,tmp_path):
    """An RFx sheet often carries an empty response template; it must not be mistaken
    for the requirement table just because it has a Qty column."""
    from openpyxl import Workbook
    wb=Workbook();ws=wb.active
    for row in [['Supplier Name','Quoted Item / Model','Qty','Rental Term Months','Price Basis / UOM'],
                ['','','100','60',''],
                [],
                ['SKU','Description','Quantity','UOM'],
                ['BX-001','5-ply RSC box',39000,'piece'],
                ['BX-002','3-ply RSC box',21000,'piece']]:
        ws.append(row)
    path=tmp_path/'rfx.xlsx';wb.save(path)
    from backend.app import intake_docs
    items,note=intake_docs.items_from_rows(intake_docs._xlsx_rows(path))
    assert [x['sku'] for x in items]==['BX-001','BX-002'],note
    assert [x['quantity'] for x in items]==[39000,21000]

def test_a_requirement_written_down_the_page_is_still_a_requirement(client,tmp_path):
    """Item / Quantity / Term as rows, not columns — the common one-line RFx layout."""
    from openpyxl import Workbook
    wb=Workbook();ws=wb.active
    for row in [['Field','Buyer Requirement'],['RFQ ID','IT-RENT-2026-0919'],
                ['Item','Apple MacBook Pro 16-inch, M4-series'],['Quantity','100 devices'],
                ['Term (months)','60'],['Delivery location','Bengaluru, Karnataka']]:
        ws.append(row)
    path=tmp_path/'vertical.xlsx';wb.save(path)
    from backend.app import intake_docs
    items,note=intake_docs.items_from_rows(intake_docs._xlsx_rows(path))
    assert len(items)==1,note
    assert items[0]['description'].startswith('Apple MacBook Pro')
    # The buyer's own word for the unit, and the term, are both kept.
    assert items[0]['quantity']==100 and items[0]['unit']=='devices' and items[0]['term_months']==60

def test_one_requirement_in_two_documents_stays_one_line(client,tmp_path,monkeypatch):
    from openpyxl import Workbook
    sheet=Workbook();ws=sheet.active
    for row in [['Item','Apple MacBook Pro 16-inch, M4-series'],['Quantity','100'],['Term (months)','60']]:
        ws.append(row)
    a=tmp_path/'rfx.xlsx';sheet.save(a)
    second=Workbook();ws2=second.active
    for row in [['Item','Apple MacBook Pro 16-inch, M4-series'],['Quantity','100 devices'],['Unit of measure','device']]:
        ws2.append(row)
    b=tmp_path/'template.xlsx';second.save(b)
    with open(a,'rb') as fa, open(b,'rb') as fb:
        value=client.post('/api/intake/attach',data={'prompt':'Both files describe the same need.'},
                          files=[('files',('rfx.xlsx',fa)),('files',('template.xlsx',fb))]).json()
    assert len(value['items'])==1,value['items']
    # The richer of the two descriptions of the same line survives.
    assert value['items'][0]['term_months']==60 and value['items'][0]['quantity']==100


def test_start_fresh_clears_the_checklist_without_deleting_history(client,monkeypatch):
    """The buyer can abandon a half-built request at any stage and begin again."""
    cheap,_=two_suppliers(client,monkeypatch)
    before=repo.read()
    assert before['rfx_status']=='approved' and before['vendors'] and before['comparison']
    assert client.get('/api/intake').json()['items']
    r=client.post('/api/intake/new',json={})
    assert r.status_code==200,r.text
    value=r.json()
    # Everything the buyer would see is empty again.
    assert value['items']==[] and value['messages']==[] and value['confirmed'] is False
    assert all(not c['value'] for c in value['checklist'])
    fresh=repo.read()
    assert fresh['rfx']['title']=='' and fresh['rfx']['items']==[]
    assert fresh['vendors']=={} and fresh['comparison']==[] and fresh['exceptions']==[]
    assert fresh['rfx_status']=='draft' and fresh['rfx']['event_id']!=before['rfx']['event_id']
    # Nothing is destroyed: the earlier snapshot and its evidence remain readable.
    assert fresh['dataset_version']>before['dataset_version']
    earlier=client.get(f"/api/datasets/{before['dataset_version']}").json()
    assert len(earlier['vendors'])==2 and earlier['rfx']['items']
    assert set(before['evidence']) <= set(fresh['evidence']), 'evidence is retained for history'

def test_start_fresh_waits_for_a_running_extraction(client,monkeypatch):
    two_suppliers(client,monkeypatch)
    monkeypatch.setattr(jobs,'ACTIVE','job-in-flight')
    assert client.post('/api/intake/new',json={}).status_code==422
    assert repo.read()['vendors'], 'nothing was cleared while extraction was running'


HTML_QUOTE_EMAIL = """From: sales@quickdeal.example
To: buyer@example.com
Subject: Quotation
MIME-Version: 1.0
Content-Type: multipart/alternative; boundary="b1"

--b1
Content-Type: text/plain; charset="utf-8"

Dear Procurement Team,
Please find our complete quote below for all four SKUs.
Freight: INR 65,000 total. Payment: Net 15.

--b1
Content-Type: text/html; charset="utf-8"

<html><body><p>Quote</p><table>
<tr><th>SKU</th><th>Offered item</th><th>Qty</th><th>Unit price</th></tr>
<tr><td>IT-LAP-001</td><td>MacBook Pro 16-inch</td><td>100</td><td>INR 202,000</td></tr>
<tr><td>IT-PRN-002</td><td>Mono laser printer</td><td>10</td><td>INR 38,500</td></tr>
</table></body></html>
--b1--
"""

def test_a_price_table_sent_as_html_is_read(client,tmp_path):
    """Suppliers routinely send the table as the HTML alternative; the plain part is
    only the covering note, and reading it alone means 'not quoted' on every line."""
    path=tmp_path/'quote.eml';path.write_text(HTML_QUOTE_EMAIL)
    text='\n'.join(b['text'] for b in pipeline.source_blocks(path))
    assert 'IT-LAP-001 | MacBook Pro 16-inch | 100 | INR 202,000' in text
    assert 'IT-PRN-002' in text and '38,500' in text
    # The covering note survives too, and is not duplicated by the HTML alternative.
    assert text.count('Freight: INR 65,000 total. Payment: Net 15.')==1
    rows=intake_docs._text_rows(path)
    assert ['IT-LAP-001','MacBook Pro 16-inch','100','INR 202,000'] in rows

@pytest.mark.parametrize('text,expected,amount',[
    ('Freight | USD 1,800 total shipment','order_total',1800),
    ('Q3: No, freight is extra at USD 1,800 for the total shipment.','order_total',1800),
    ('Freight: INR 65,000 total. Payment: Net 15.','order_total',65000),
    ('Delivery charges Rs. 12,500 one-time','order_total',12500),
    ('Freight ₹0.42 per piece','flat',None),
    ('Freight 2% of material value','percent',None),
])
def test_freight_quoted_for_the_whole_order_is_understood(text,expected,amount):
    terms=pipeline.terms_from_text(text)
    assert terms['freight']==expected,terms
    if amount:assert terms['freight_total']==amount

def test_whole_order_freight_is_apportioned_by_line_value(client,monkeypatch):
    complete_intake(client,items=[{'description':'Laptop','quantity':100,'unit':'unit'},
                                  {'description':'Printer','quantity':10,'unit':'unit'}])
    vid=client.post('/api/proposals',data={'supplier_name':'GlobalSource','email':'g@example.com',
        'body':'Freight: INR 65,000 total for the shipment.'}).json()['vendor_id']
    def extract(*args,**kwargs):
        return {'facts':[priced(1,raw_price=202000,raw_currency='INR',raw_basis='piece',covers_units=1,covers_period_months=None),
                         priced(2,raw_price=38500,raw_currency='INR',raw_basis='piece',covers_units=1,covers_period_months=None)],
                'terms':pipeline.terms_from_text('Freight: INR 65,000 total for the shipment.'),
                'questionnaire':{'Q1':{'answer':'Yes','confidence':.99}},'quality':'pass','method':'mock',
                'unresolved':[],'blocks':[],'model':None,'captured_terms':[]}
    monkeypatch.setattr(pipeline,'extract',extract)
    graph.run('process')
    d=repo.read()
    first,second=(r['vendors'][vid] for r in d['comparison'])
    assert first['landed_unit_cost'] is not None and second['landed_unit_cost'] is not None
    # Every line carries the same proportion of the freight, and the parts add up.
    share_one=first['landed_unit_cost']-202000
    share_two=second['landed_unit_cost']-38500
    assert abs(share_one/202000-share_two/38500)<1e-6
    assert abs(share_one*100+share_two*10-65000)<1, 'apportioned freight must total the quoted charge'
    assert any('apportioned across lines by value' in s for s in first['transformation_steps'])


def test_a_stated_period_is_ranked_against_the_other_responses(client,monkeypatch):
    """The buyer wants where each supplier stands, not four sentences to read."""
    complete_intake(client)
    names={}
    for name,warranty in [('LongCover','5-year onsite warranty'),('MidCover','3 years'),
                          ('ShortCover','12-month reseller warranty'),('VagueCover','Standard OEM warranty')]:
        names[name]=client.post('/api/proposals',data={'supplier_name':name,'email':f'{name}@example.com',
                                                       'body':'Freight included.'}).json()['vendor_id']
    warranties={v:w for (n,w),v in zip([('LongCover','5-year onsite warranty'),('MidCover','3 years'),
                                        ('ShortCover','12-month reseller warranty'),('VagueCover','Standard OEM warranty')],names.values())}
    def extract(vendor_id,*args,**kwargs):
        return {'facts':[priced(raw_basis='piece',raw_price=100,raw_currency='INR')],
                'terms':pipeline.terms_from_text('Freight included.'),
                'questionnaire':{'Q1':{'answer':'Yes','confidence':.99}},'quality':'pass','method':'mock',
                'unresolved':[],'blocks':[],'model':None,
                'captured_terms':[{'term_type':'warranty','value':warranties[vendor_id],'confidence':.9,'source_text':'w'}]}
    monkeypatch.setattr(pipeline,'extract',extract)
    graph.run('process')
    standing={v['name']:v['terms']['warranty']['standing'] for v in insights.overview()['suppliers']}
    assert standing['LongCover']=='Strongest · 5 years'
    assert standing['MidCover']=='Middle · 3 years'
    assert standing['ShortCover']=='Weakest · 1 year'
    # A supplier who stated a warranty without a period is not ranked against those who did.
    assert standing['VagueCover']=='Stated, no period given'

def test_durations_are_read_the_way_suppliers_write_them():
    months=insights.duration_months
    assert months('5-year onsite warranty')==60 and months('3 years')==36
    assert months('12-month reseller warranty')==12 and months('18 months')==18
    assert months('Standard OEM warranty') is None
    assert months('2 weeks')==pytest.approx(0.46,abs=.01)
    # "Net 15" is a payment period even though it names no unit.
    assert months('Payment: Net 15',allow_net=True)==pytest.approx(0.49,abs=.01)
    assert months('Payment: Net 15') is None


def four_line_event(client,monkeypatch):
    """Four lines, two suppliers, each cheapest on a different pair."""
    complete_intake(client,items=[{'description':f'Item {n}','quantity':10*n,'unit':'unit','sku':f'SKU-{n:03d}'}
                                  for n in (1,2,3,4)])
    cheap=client.post('/api/proposals',data={'supplier_name':'AlphaCo','email':'a@example.com','body':'Freight included.'}).json()['vendor_id']
    other=client.post('/api/proposals',data={'supplier_name':'BetaCo','email':'b@example.com','body':'Freight included.'}).json()['vendor_id']
    prices={cheap:[100,900,300,700],other:[200,800,400,600]}
    leads={cheap:[6,7,8,9],other:[9,8,7,6]}
    def extract(vendor_id,*args,**kwargs):
        return {'facts':[priced(n,raw_price=prices[vendor_id][n-1],raw_currency='INR',raw_basis='piece',
                                covers_units=1,covers_period_months=None,lead_days=leads[vendor_id][n-1]) for n in (1,2,3,4)],
                'terms':pipeline.terms_from_text('Freight included.'),
                'questionnaire':{'Q1':{'answer':'Yes','confidence':.99}},'quality':'pass','method':'mock',
                'unresolved':[],'blocks':[],'model':None,'captured_terms':[]}
    monkeypatch.setattr(pipeline,'extract',extract)
    graph.run('process')
    return cheap,other

def test_a_question_about_skus_is_answered_per_sku(client,monkeypatch):
    """"SKU-wise cost per vendor" used to collapse into one total per supplier."""
    cheap,other=four_line_event(client,monkeypatch)
    result=knowledge.by_line({})
    assert result['axis']=='requirement' and len(result['rows'])==4
    assert [r['label'].split(' · ')[0] for r in result['rows']]==['SKU-001','SKU-002','SKU-003','SKU-004']
    # The cheapest supplier is marked on each line independently.
    assert [r['winners'] for r in result['rows']]==[[f'supplier:{cheap}'],[f'supplier:{other}'],
                                                    [f'supplier:{cheap}'],[f'supplier:{other}']]
    assert result['rows'][0]['cells'][f'supplier:{cheap}']['numeric']==100
    assert result['rows'][0]['cells'][f'supplier:{cheap}']['evidence_id'], 'every figure keeps its source'
    verdict=knowledge.line_verdict(result)
    assert 'cheapest on 2' in verdict and len(verdict.split())<=20

def test_the_per_sku_view_can_answer_on_lead_time_too(client,monkeypatch):
    four_line_event(client,monkeypatch)
    result=knowledge.by_line({'dimensions':['delivery']})
    assert result['metric']=='delivery'
    assert result['rows'][0]['cells'][result['suppliers'][0]['id']]['display']=='6 days'
    assert 'fastest on' in knowledge.line_verdict(result)

def test_the_per_sku_view_narrows_to_the_lines_asked_about(client,monkeypatch):
    cheap,_=four_line_event(client,monkeypatch)
    result=knowledge.by_line({'requirements':[2,3]})
    assert [r['key'] for r in result['rows']]==['line:2','line:3']
    one=knowledge.by_line({'requirements':[1],'suppliers':[f'supplier:{cheap}']})
    assert len(one['rows'])==1 and [s['id'] for s in one['suppliers']]==[f'supplier:{cheap}']

def test_a_line_without_a_usable_price_says_why(client,monkeypatch):
    cheap,other=four_line_event(client,monkeypatch)
    repo.change('withdraw','test',lambda d:(d['vendors'][other]['facts'].pop(0),pipeline.rebuild(d)))
    row=knowledge.by_line({})['rows'][0]
    assert row['cells'][f'supplier:{other}']['display']=='Not quoted'
    assert row['winners']==[f'supplier:{cheap}'], 'a missing price never wins by default'

def test_the_graph_query_endpoint_exposes_the_same_breakdown(client,monkeypatch):
    four_line_event(client,monkeypatch)
    answer=client.post('/api/graph/query',json={'breakdown':'requirement'}).json()
    assert answer['axis']=='requirement' and len(answer['rows'])==4
    assert 'cheapest on' in answer['verdict']
    supplier_level=client.post('/api/graph/query',json={'breakdown':'supplier','dimensions':['cost']}).json()
    assert supplier_level['rows'][0]['key']=='cost'


def test_two_award_strategies_are_compared_side_by_side(client,monkeypatch):
    """"Everything to one supplier vs a split" is two awards, not one."""
    from backend.app import main
    cheap,other=four_line_event(client,monkeypatch)
    strategy=lambda label,excluded:{'label':label,'required_supplier_ids':[],'excluded_supplier_ids':excluded,
                                    'qualified_suppliers_only':True,'max_supplier_spend_share':None,'max_delivery_days':None}
    answer=main.compare_awards({'headline':'Award strategies compared','scenarios':[
        strategy('All to AlphaCo',[f'supplier:{other}']),strategy('Split by best price',[])]},'q')
    assert answer['kind']=='scenarios' and len(answer['runs'])==2
    whole,split=answer['runs']
    # AlphaCo alone: 100x10 + 900x20 + 300x30 + 700x40. Split takes the cheaper of each line.
    assert whole['award_total_inr']==56000 and split['award_total_inr']==50000
    assert [m['name'] for m in whole['supplier_mix']]==['AlphaCo']
    assert {m['name'] for m in split['supplier_mix']}=={'AlphaCo','BetaCo'}
    assert answer['best']=='Split by best price'
    assert '6,000' in answer['text'].replace('0.1 lakh','6,000') or 'Split by best price' in answer['text']
    assert len(answer['text'].split())<=30

def test_an_infeasible_strategy_is_reported_not_hidden(client,monkeypatch):
    from backend.app import main
    cheap,other=four_line_event(client,monkeypatch)
    answer=main.compare_awards({'headline':'x','scenarios':[
        {'label':'Everyone','required_supplier_ids':[],'excluded_supplier_ids':[],'qualified_suppliers_only':True,
         'max_supplier_spend_share':None,'max_delivery_days':None},
        {'label':'Impossible one-day delivery','required_supplier_ids':[],'excluded_supplier_ids':[],
         'qualified_suppliers_only':True,'max_supplier_spend_share':None,'max_delivery_days':1}]},'q')
    feasible,impossible=answer['runs']
    assert feasible['award_total_inr']==50000
    assert impossible['award_total_inr'] is None and impossible['status']=='infeasible'
    assert answer['best']=='Everyone'
    # The comparison still reports a usable number rather than failing outright.
    assert 'Everyone' in answer['text']


def test_an_award_can_be_broken_down_line_by_line(client,monkeypatch):
    """"Which SKUs go to whom" is answerable from the award itself."""
    from backend.app import main
    cheap,other=four_line_event(client,monkeypatch)
    answer=main.award_allocation({'headline':'Split allocation','scenarios':[
        {'label':'Split by best price','required_supplier_ids':[],'excluded_supplier_ids':[],
         'qualified_suppliers_only':True,'max_supplier_spend_share':None,'max_delivery_days':None}]},'q')
    assert answer['kind']=='matrix' and answer['axis']=='allocation'
    assert [r['winners'] for r in answer['rows']]==[[cheap],[other],[cheap],[other]]
    # The winner's cell carries the line value and how it was reached; the rest are blank.
    won=answer['rows'][0]['cells'][cheap]
    assert won['numeric']==1000 and won['detail']=='10 × ₹100.00' and won['evidence_id']
    assert answer['rows'][0]['cells'][other]['display']=='—'
    assert 'AlphaCo: 2 lines' in answer['text'] and 'BetaCo: 2 lines' in answer['text']
    assert len(answer['text'].split())<=25

def test_a_sole_award_gives_every_line_to_one_supplier(client,monkeypatch):
    from backend.app import main
    cheap,other=four_line_event(client,monkeypatch)
    answer=main.award_allocation({'headline':'All to AlphaCo','scenarios':[
        {'label':'All to AlphaCo','required_supplier_ids':[],'excluded_supplier_ids':[f'supplier:{other}'],
         'qualified_suppliers_only':True,'max_supplier_spend_share':None,'max_delivery_days':None}]},'q')
    assert [s['id'] for s in answer['suppliers']]==[cheap]
    assert all(r['winners']==[cheap] for r in answer['rows'])
    assert 'AlphaCo: 4 lines' in answer['text']

def test_an_allocation_that_cannot_be_solved_says_so(client,monkeypatch):
    from backend.app import main
    four_line_event(client,monkeypatch)
    answer=main.award_allocation({'headline':'x','scenarios':[
        {'label':'One-day delivery','required_supplier_ids':[],'excluded_supplier_ids':[],
         'qualified_suppliers_only':True,'max_supplier_spend_share':None,'max_delivery_days':1}]},'q')
    assert answer['kind']=='text' and answer['title']=='No feasible award'

def test_the_analysis_endpoint_accepts_the_conversation_so_far(client,monkeypatch):
    """A follow-up carries the earlier turns; they are passed to the router, not ignored."""
    four_line_event(client,monkeypatch)
    seen={}
    def interpret(question,schema,history=None):
        seen['history']=history
        return None
    monkeypatch.setattr(insights,'interpret',interpret)
    r=client.post('/api/analysis',json={'question':'which SKUs go to each?','history':[
        {'question':'compare the split','summary':'Split GlobalSource & QuickDeal: ₹23,125,870'}]})
    assert r.status_code==200
    assert seen['history']==[{'question':'compare the split','summary':'Split GlobalSource & QuickDeal: ₹23,125,870'}]
    # Oversized histories are rejected rather than silently truncated.
    assert client.post('/api/analysis',json={'question':'x','history':[{'question':'q','summary':'s'}]*9}).status_code==422


def test_an_award_can_be_split_inside_a_line(client,monkeypatch):
    """"25% of each SKU to each of four vendors" needs quantity-level allocation."""
    from backend.app import main
    cheap,other=four_line_event(client,monkeypatch)
    answer=main.award_allocation({'headline':'Equal split','scenarios':[
        {'label':'50/50 on every line','required_supplier_ids':[],'excluded_supplier_ids':[],
         'qualified_suppliers_only':True,'max_supplier_spend_share':None,'max_delivery_days':None,
         'allocation_granularity':'quantity','equal_split':True}]},'q')
    assert answer['kind']=='matrix' and answer['axis']=='allocation'
    # Every line is divided, so no single supplier "wins" it.
    assert all(r['winners']==[] for r in answer['rows'])
    both=answer['rows'][0]['cells']
    assert both[cheap]['numeric']==500 and both[other]['numeric']==1000, 'half of 10 units at each price'
    assert both[cheap]['detail'].startswith('50% ·')
    # Equal quantity shares, unequal spend, and the parts still add up.
    total=sum(c['numeric'] for r in answer['rows'] for c in r['cells'].values() if c['numeric'])
    # Half of each line from each supplier: the mean of the two sole awards (56,000 and 54,000).
    assert round(total)==55000 and '% of spend' in answer['text']

def test_an_equal_split_needs_divisible_lines(client,monkeypatch):
    from backend.app.scenario import solve_scenario
    four_line_event(client,monkeypatch)
    result=solve_scenario({'equal_split':True},repo.read())
    assert result['status']=='infeasible' and 'quantity-level allocation' in result['reason']

def test_whole_line_awards_are_unchanged_by_divisibility(client,monkeypatch):
    from backend.app.scenario import solve_scenario
    cheap,other=four_line_event(client,monkeypatch)
    whole=solve_scenario({},repo.read())
    assert whole['award_total_inr']==50000 and all(a['line_share']==1 for a in whole['allocation'])
    # Allowing division cannot cost more than the best whole-line award.
    divisible=solve_scenario({'allocation_granularity':'quantity'},repo.read())
    assert divisible['award_total_inr']<=whole['award_total_inr']+1e-6

def test_the_knowledge_base_is_written_from_the_graph(client,monkeypatch):
    from backend.app import dossier
    four_line_event(client,monkeypatch)
    files=dossier.build()
    assert set(files)>={'index.md','requirement.md','supplier-alphaco.md','supplier-betaco.md'}
    alpha=files['supplier-alphaco.md']
    assert alpha.startswith('# AlphaCo')
    # The dossier states the same numbers the structured query returns.
    assert '| Qualification | pass |' in alpha and 'SKU-001' in alpha
    assert '## Quoted lines' in alpha and '## How each price was normalized' in alpha
    assert '## Answers to the buyer’s questions' in alpha
    requirement=files['requirement.md']
    assert 'SKU-004' in requirement and '## Questions asked of every supplier' in requirement
    # And it lands on disk, versioned with the snapshot it came from.
    folder=dossier.write()
    assert folder.name==f"v{repo.read()['dataset_version']}"
    assert (folder/'supplier-alphaco.md').read_text()==alpha

def test_the_knowledge_base_endpoints_serve_markdown(client,monkeypatch):
    four_line_event(client,monkeypatch)
    listing=client.get('/api/knowledge').json()
    assert listing['dataset_version']==repo.read()['dataset_version']
    assert any(d['name']=='supplier-alphaco.md' for d in listing['documents'])
    page=client.get('/api/knowledge/supplier-alphaco.md')
    assert page.status_code==200 and page.headers['content-type'].startswith('text/markdown')
    assert client.get('/api/knowledge/supplier-nobody.md').status_code==404

def test_the_conversation_is_logged_on_the_server(client,monkeypatch):
    """A reload must not lose the thread, so the log lives with the event."""
    four_line_event(client,monkeypatch)
    seen={}
    def interpret(question,schema,history=None):
        seen[question]=history
        return None
    monkeypatch.setattr(insights,'interpret',interpret)
    client.post('/api/analysis',json={'question':'cheapest overall?','session_id':'s-1'})
    client.post('/api/analysis',json={'question':'and per SKU?','session_id':'s-1'})
    turns=client.get('/api/sessions/s-1').json()['turns']
    assert [t['question'] for t in turns]==['cheapest overall?','and per SKU?']
    assert all(t['at'] for t in turns)
    # The second question was routed with the first turn as context.
    assert seen['and per SKU?'][0]['question']=='cheapest overall?'
    # Separate chats keep separate memories.
    client.post('/api/analysis',json={'question':'fresh start','session_id':'s-2'})
    assert len(client.get('/api/sessions/s-2').json()['turns'])==1
    assert {s['id'] for s in client.get('/api/sessions').json()}>={'s-1','s-2'}


def test_the_buyers_own_percentages_are_honoured_on_every_line(client,monkeypatch):
    """60/40 means 60/40 on each SKU — not whatever the optimizer prefers."""
    from backend.app.scenario import solve_scenario
    cheap,other=four_line_event(client,monkeypatch)
    result=solve_scenario({'allocation_granularity':'quantity',
                           'supplier_line_shares':{cheap:0.6,other:0.4}},repo.read())
    assert result['status']=='ok'
    for line in (1,2,3,4):
        parts={a['vendor_id']:a['line_share'] for a in result['allocation'] if a['line_no']==line}
        assert parts=={cheap:0.6,other:0.4}, f'line {line} was {parts}'
    # Quantities follow the share, and the total is the weighted blend of both quotes.
    first={a['vendor_id']:a for a in result['allocation'] if a['line_no']==1}
    assert first[cheap]['qty']==6 and first[other]['qty']==4
    assert result['award_total_inr']==round(.6*56000+.4*54000,2)

def test_a_partial_share_leaves_the_rest_to_the_optimizer(client,monkeypatch):
    from backend.app.scenario import solve_scenario
    cheap,other=four_line_event(client,monkeypatch)
    result=solve_scenario({'allocation_granularity':'quantity','supplier_line_shares':{cheap:0.25}},repo.read())
    for line in (1,2,3,4):
        parts={a['vendor_id']:a['line_share'] for a in result['allocation'] if a['line_no']==line}
        assert abs(parts[cheap]-0.25)<1e-6
        assert abs(sum(parts.values())-1)<1e-6, 'the remainder is still allocated'

def test_shares_that_cannot_be_honoured_say_why(client,monkeypatch):
    from backend.app.scenario import solve_scenario
    cheap,other=four_line_event(client,monkeypatch)
    repo.change('withdraw','test',lambda d:(d['vendors'][other]['facts'].pop(0),pipeline.rebuild(d)))
    result=solve_scenario({'allocation_granularity':'quantity',
                           'supplier_line_shares':{cheap:0.6,other:0.4}},repo.read())
    assert result['status']=='infeasible'
    assert 'no usable price on line 1' in result['reason'] and 'BetaCo' in result['reason']
    # And the same shares on whole lines are refused rather than silently rounded.
    assert 'quantity-level allocation' in solve_scenario(
        {'supplier_line_shares':{cheap:0.6,other:0.4}},repo.read())['reason']

def test_a_supplier_the_buyer_names_is_never_filtered_out(client,monkeypatch):
    """Qualification gates the default award, not a supplier chosen by hand."""
    from backend.app.scenario import solve_scenario
    cheap,other=four_line_event(client,monkeypatch)
    repo.change('pending','test',lambda d:d['vendors'][other].update(quality='pending') or pipeline.rebuild(d))
    default=solve_scenario({},repo.read())
    assert other not in {m['vendor_id'] for m in default['vendor_mix']}, 'unqualified by default'
    named=solve_scenario({'allocation_granularity':'quantity',
                          'supplier_line_shares':{cheap:0.5,other:0.5}},repo.read())
    assert named['status']=='ok'
    assert {m['vendor_id'] for m in named['vendor_mix']}=={cheap,other}

def test_comparing_two_splits_returns_both_with_their_lines(client,monkeypatch):
    from backend.app import main
    cheap,other=four_line_event(client,monkeypatch)
    split=lambda label,a,b:{'label':label,'required_supplier_ids':[],'excluded_supplier_ids':[],
                            'qualified_suppliers_only':True,'max_supplier_spend_share':None,'max_delivery_days':None,
                            'allocation_granularity':'quantity','equal_split':False,
                            'supplier_line_shares':[{'supplier_id':f'supplier:{cheap}','share':a},
                                                    {'supplier_id':f'supplier:{other}','share':b}]}
    answer=main.compare_awards({'headline':'60/40 vs 50/50','scenarios':[
        split('60/40',0.6,0.4),split('50/50',0.5,0.5)]},'q',allocations=True)
    assert [r['label'] for r in answer['runs']]==['60/40','50/50']
    assert all(len(r['allocation'])==8 for r in answer['runs']), 'four lines split two ways'
    assert {a['share'] for a in answer['runs'][0]['allocation']}=={0.6,0.4}
    assert {a['share'] for a in answer['runs'][1]['allocation']}=={0.5}
    # A difference of a few hundred rupees is reported in rupees, not "0.0 lakh".
    assert 'lakh' not in answer['text'] and '₹' in answer['text']


FREIGHT_WORDING = [
    # A quoted amount with no per-unit qualifier is a charge for the whole order,
    # whether or not the supplier wrote the word "total".
    ('Commercial summary: subtotal $207,610.00; international + India delivery freight $1,200.00; taxes extra.',
     'order_total', 1200.0, 'USD'),
    ('Q3: No, freight is extra: international + India delivery freight $1,200.00.', 'order_total', 1200.0, 'USD'),
    ('Freight: INR 65,000 total. Payment: Net 15.', 'order_total', 65000.0, 'INR'),
    ('Freight | USD 1,800 total shipment', 'order_total', 1800.0, 'USD'),
    ('Freight ₹0.42 per piece', 'flat', None, None),
    ('Freight 2% of material value', 'percent', None, None),
    # An amount elsewhere in the sentence never overrides a stated inclusion.
    ('Prices are inclusive of delivery to site; unit price USD 250', 'included', None, None),
    ('All rates include freight to Bengaluru. Subtotal USD 90,000.', 'included', None, None),
    ('Ex-works Pune', 'excluded', None, None),
    ('Lead time 21 days', 'unknown', None, None),
]

@pytest.mark.parametrize('text,basis,amount,currency',FREIGHT_WORDING)
def test_freight_amounts_are_read_the_way_suppliers_write_them(text,basis,amount,currency):
    terms=pipeline.terms_from_text(text)
    assert terms['freight']==basis,terms
    if amount is not None:
        assert terms['freight_total']==amount and terms['freight_currency']==currency

def test_a_dollar_quote_with_whole_order_freight_still_lands(client,monkeypatch):
    """Currency conversion and order-level freight have to work together."""
    complete_intake(client,items=[{'description':'Laptop','quantity':100,'unit':'unit'},
                                  {'description':'Dock','quantity':50,'unit':'unit'}])
    vid=client.post('/api/proposals',data={'supplier_name':'GlobalTech','email':'g@example.com',
        'body':'Merchandise subtotal $207,610.00; international + India delivery freight $1,200.00.'}).json()['vendor_id']
    def extract(*args,**kwargs):
        return {'facts':[priced(1,raw_price=1780,raw_currency='USD',raw_basis='piece',covers_units=1,covers_period_months=None),
                         priced(2,raw_price=820,raw_currency='USD',raw_basis='piece',covers_units=1,covers_period_months=None)],
                'terms':pipeline.terms_from_text('Merchandise subtotal $207,610.00; international + India delivery freight $1,200.00.'),
                'questionnaire':{'Q1':{'answer':'Yes','confidence':.99}},'quality':'pass','method':'mock',
                'unresolved':[],'blocks':[],'model':None,'captured_terms':[]}
    monkeypatch.setattr(pipeline,'extract',extract)
    graph.run('process')
    fx=repo.read()['rfx']['commercial_rules']['fx_rate_usd_inr']
    first,second=(r['vendors'][vid] for r in repo.read()['comparison'])
    assert first['unit_price']==round(1780*fx,2), 'converted once, at the RFx rate'
    assert first['landed_unit_cost'] is not None and second['landed_unit_cost'] is not None
    # The USD freight is converted too, then shared out by line value.
    freight=(first['landed_unit_cost']-first['unit_price'])*100+(second['landed_unit_cost']-second['unit_price'])*50
    assert abs(freight-1200*fx)<1
    assert any('apportioned across lines by value' in step for step in first['transformation_steps'])


def test_a_question_is_classified_by_its_topic_not_its_wording(client):
    """The compliance question quotes the buyer's requirement, which may mention GST,
    freight or warranty. That must not make it the answer to those questions."""
    from backend.app import knowledge
    complete_intake(client)
    asked={q['id']:q for q in client.get('/api/rfx').json()['questionnaire']}
    assert asked['Q1']['topic']=='compliance' and asked['Q7']['topic']=='tax'
    assert knowledge.question_kind(asked['Q1']) is None, 'compliance is not a commercial term'
    assert knowledge.question_kind(asked['Q7'])=='tax'
    assert knowledge.question_kind(asked['Q3'])=='freight'
    assert knowledge.question_kind(asked['Q5'])=='payment'
    assert knowledge.question_kind(asked['Q8'])=='warranty'
    # A questionnaire without topics still reads its own stem, not the quoted requirement.
    legacy={'id':'Q1','question':'Confirm your offer complies with these mandatory requirements: '
                                 'valid GST registration, ISO 9001 and freight to site'}
    assert knowledge.question_kind(legacy) is None
    assert knowledge.question_kind({'id':'Q7','question':'State the currency and tax treatment of your '
                                                        'prices. Buyer basis: INR'})=='tax'

def test_the_tax_row_reads_the_tax_answer(client,monkeypatch):
    complete_intake(client)
    client.post('/api/proposals',data={'supplier_name':'Apex','email':'a@example.com','body':'Freight included.'})
    def extract(*args,**kwargs):
        return {'facts':[priced(raw_basis='piece',raw_price=1000,raw_currency='INR')],
                'terms':pipeline.terms_from_text('Freight included.'),
                'questionnaire':{'Q1':{'answer':'Yes. ISO 9001 valid, GST registration valid.','confidence':.99},
                                 'Q5':{'answer':'Net 45 days','confidence':.9},
                                 'Q7':{'answer':'Currency INR; 18% GST extra.','confidence':.9}},
                'quality':'pass','method':'mock','unresolved':[],'blocks':[],'model':None,'captured_terms':[]}
    monkeypatch.setattr(pipeline,'extract',extract)
    graph.run('process')
    terms=insights.overview()['suppliers'][0]['terms']
    assert terms['tax']['value']=='Currency INR; 18% GST extra.' and terms['tax']['source']=='Answer to Q7'
    assert 'ISO 9001' not in terms['tax']['value']
    # 45 days of credit is shown as days, not rounded to a month.
    assert terms['payment']['standing'].endswith('45 days')

def test_the_same_period_written_differently_ranks_the_same(client,monkeypatch):
    complete_intake(client)
    names={}
    for name,payment in [('MonthCo','Net 1 month'),('DayCo','Net 30 days'),('SlowCo','Net 60 days')]:
        names[client.post('/api/proposals',data={'supplier_name':name,'email':f'{name}@e.com',
                          'body':'Freight included.'}).json()['vendor_id']]=payment
    def extract(vendor_id,*args,**kwargs):
        return {'facts':[priced(raw_basis='piece',raw_price=1000,raw_currency='INR')],
                'terms':pipeline.terms_from_text('Freight included.'),
                'questionnaire':{'Q1':{'answer':'Yes','confidence':.99},
                                 'Q5':{'answer':names[vendor_id],'confidence':.9}},
                'quality':'pass','method':'mock','unresolved':[],'blocks':[],'model':None,'captured_terms':[]}
    monkeypatch.setattr(pipeline,'extract',extract)
    graph.run('process')
    standing={v['name']:v['terms']['payment']['standing'] for v in insights.overview()['suppliers']}
    assert standing['SlowCo'].startswith('Strongest')
    # "1 month" and "30 days" are the same commercial term, so neither outranks the other.
    assert standing['MonthCo']==standing['DayCo'], standing
    assert standing['MonthCo'].startswith('Weakest')


COMPLIANCE_ANSWERS = [
    ('Yes. ISO 9001: Valid - CERT-2028, GST registration: Valid, Dell authorization: Valid.', 'pass'),
    # A list of valid certificates is a confirmation even without the word "yes".
    ('ISO 9001:2015: Valid - SGC-77192 (2027); Dell, Lenovo and Cisco authorized channels: Valid '
     '(OEM auth attached); GST: registered.', 'pass'),
    # A supplier who says plainly that something is missing has failed, not gone unread.
    ('Non-compliant with Cisco OEM authorization requirement (Cisco NOT ENCLOSED).', 'fail'),
    ('Non-compliant / Not qualified on mandatory ISO 9001 (EXPIRED 31 Aug 2026).', 'fail'),
    ('No.', 'fail'),
    ('Our ISO 9001 renewal is in progress.', 'pending'),
    ('', 'pending'),
]

@pytest.mark.parametrize('answer,verdict',COMPLIANCE_ANSWERS)
def test_compliance_is_read_not_pattern_matched(answer,verdict):
    assert pipeline.read_verdict(answer)[0]==verdict,pipeline.read_verdict(answer)

def test_qualification_states_which_item_decided_it(client,monkeypatch):
    complete_intake(client)
    vid=client.post('/api/proposals',data={'supplier_name':'Metro','email':'m@example.com','body':'Freight included.'}).json()['vendor_id']
    def extract(*args,**kwargs):
        return {'facts':[priced(raw_basis='piece',raw_price=1000,raw_currency='INR')],
                'terms':pipeline.terms_from_text('Freight included.'),
                'questionnaire':{'Q1':{'answer':'Non-compliant with the Cisco OEM authorization requirement.','confidence':.95}},
                'quality':'pass','method':'mock','unresolved':[],'blocks':[],'model':None,'captured_terms':[]}
    monkeypatch.setattr(pipeline,'extract',extract)
    graph.run('process')
    vendor=repo.read()['vendors'][vid]
    assert vendor['quality']=='fail', 'a stated non-compliance is a failure'
    assert vendor['quality_reason']
    review=next(e for e in repo.read()['exceptions'] if e['type']=='qualification')
    assert vendor['quality_reason'].rstrip('.') in review['detail']

def test_the_buyer_can_record_a_qualification_decision(client,monkeypatch):
    """The buyer is the authority on whether a declaration satisfies their requirement."""
    complete_intake(client)
    vid=client.post('/api/proposals',data={'supplier_name':'GlobalTech','email':'g@example.com','body':'Freight included.'}).json()['vendor_id']
    def extract(*args,**kwargs):
        return {'facts':[priced(raw_basis='piece',raw_price=1000,raw_currency='INR')],
                'terms':pipeline.terms_from_text('Freight included.'),
                'questionnaire':{'Q1':{'answer':'We supply Dell, Lenovo and Cisco hardware across India.','confidence':.95}},
                'quality':'pass','method':'mock','unresolved':[],'blocks':[],'model':None,'captured_terms':[]}
    monkeypatch.setattr(pipeline,'extract',extract)
    graph.run('process')
    assert repo.read()['vendors'][vid]['quality']=='pending'
    assert solve_scenario({},repo.read())['status']=='infeasible', 'unqualified by default'
    review=next(e for e in repo.read()['exceptions'] if e['type']=='qualification')
    r=client.post('/api/exceptions/'+review['id']+'/resolve',json={'action':'qualify','actor':'Priya Menon',
        'reason':'GST certificate verified on the portal against their India entity.',
        'expected_version':repo.read()['dataset_version']})
    assert r.status_code==200,r.text
    vendor=repo.read()['vendors'][vid]
    assert vendor['quality']=='pass' and vendor['quality_override']['actor']=='Priya Menon'
    assert 'verified on the portal' in vendor['quality_reason']
    assert not [e for e in repo.read()['exceptions'] if e['type']=='qualification' and e['status']!='resolved_by_buyer']
    assert solve_scenario({},repo.read())['status']=='ok'
    # The decision belongs to a qualification review and nowhere else.
    other=next((e for e in repo.read()['exceptions'] if e['type']!='qualification' and e['status']!='resolved_by_buyer'),None)
    if other:
        assert client.post('/api/exceptions/'+other['id']+'/resolve',json={'action':'qualify','actor':'Priya Menon',
            'reason':'Not applicable here.','expected_version':repo.read()['dataset_version']}).status_code==422
